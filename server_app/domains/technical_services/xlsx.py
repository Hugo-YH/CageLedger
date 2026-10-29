"""Service statement worksheet with quantities kept in their original units."""

from openpyxl.styles import Alignment, Font, PatternFill

from server_app.shared.xlsx import set_literal_cell


def _append_business_row(sheet, values):
    sheet.append(values)
    for column, value in enumerate(values, start=1):
        set_literal_cell(sheet.cell(sheet.max_row, column), value)


def write_statement_sheet(sheet, statement):
    sheet.append(["实验动物中心技术服务费结算单"])
    sheet.merge_cells("A1:H1")
    _append_business_row(sheet, ["月份", statement["month"], "项目负责人", statement["pi"]])
    _append_business_row(
        sheet, ["服务日数", statement.get("serviceDayCount", 0), "技术服务费（元）", statement["totalAmount"]]
    )
    sheet.append(["IACUC", "服务设施", "服务项目", "服务日期", "单位", "计费数量", "单价（元）", "金额（元）"])
    for row in statement.get("technicalServiceRows", []):
        _append_business_row(
            sheet,
            [
                row["iacuc"],
                "珠江新城" if row["facility"] == "zhujiang" else "生物岛",
                row["serviceName"],
                row["date"],
                row["unit"],
                row["quantity"],
                row["unitPrice"],
                row["amount"],
            ],
        )
    last = sheet.max_row
    sheet.append(["合计", "", "", "", "", "", "", f"=SUM(H5:H{last})" if last >= 5 else 0])
    sheet.append(["项目负责人签字：", "", "", "日期："])
    for row in sheet:
        for cell in row:
            cell.font = Font(name="Microsoft YaHei", size=11, bold=cell.row in (1, 4, last + 1))
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for cell in sheet[4]:
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in sheet.iter_rows(min_row=5, max_row=last + 1, min_col=7, max_col=8):
        for cell in row:
            cell.number_format = "0.00"
    sheet["D3"].number_format = "0.00"
    for column, width in zip("ABCDEFGH", (16, 16, 28, 16, 12, 14, 16, 18), strict=True):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A5"
    sheet.print_title_rows = "1:4"
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
