"""Generate Word by filling the user's retained laboratory forms in place."""

import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from lxml import etree

from .attachment_metadata import project_ids
from .document_layout import (
    cell_text,
    copy_paragraph,
    horizontal_result,
    image_table,
    paragraph_text,
    reagent_table,
    replace_element,
    sample_table,
)
from .document_patterns import result_patterns, result_prefix, sized_result_pattern, template_name

TEMPLATES = Path(__file__).resolve().parents[2] / "resources" / "quarantine" / "templates"

WORDPROCESSINGML = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
VML = "urn:schemas-microsoft-com:vml"
OFFICE = "urn:schemas-microsoft-com:office:office"


def normalize_style_order(content):
    """Repair the known legacy style-order violation without changing style values."""
    root = etree.fromstring(content)
    style_tag = f"{{{WORDPROCESSINGML}}}style"
    ui_priority_tag = f"{{{WORDPROCESSINGML}}}uiPriority"
    quick_format_tag = f"{{{WORDPROCESSINGML}}}qFormat"
    changed = False
    for style in root.findall(style_tag):
        ui_priority = style.find(ui_priority_tag)
        quick_format = style.find(quick_format_tag)
        if ui_priority is None:
            continue
        first_after_ui_priority = min(
            (
                style.index(element)
                for element in (
                    style.find(f"{{{WORDPROCESSINGML}}}semiHidden"),
                    style.find(f"{{{WORDPROCESSINGML}}}unhideWhenUsed"),
                    quick_format,
                )
                if element is not None
            ),
            default=len(style),
        )
        if style.index(ui_priority) > first_after_ui_priority:
            style.remove(ui_priority)
            style.insert(first_after_ui_priority, ui_priority)
            changed = True
    if not changed:
        return content
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)


def normalize_footer_shape_ids(content):
    """Give only duplicate legacy VML shapes a distinct ID in exported footers."""
    root = etree.fromstring(content)
    shape_tag = f"{{{VML}}}shape"
    office_spid = f"{{{OFFICE}}}spid"
    reserved = {shape.get("id") for shape in root.iter(shape_tag) if shape.get("id")}
    seen, changed = set(), False
    for shape in root.iter(shape_tag):
        shape_id = shape.get("id")
        if not shape_id or shape_id not in seen:
            if shape_id:
                seen.add(shape_id)
            continue
        prefix, separator, suffix = shape_id.rpartition("s")
        if not separator or not suffix.isdigit():
            raise ValueError(f"Unsupported duplicate VML shape ID: {shape_id}")
        next_number = int(suffix) + 1
        replacement = f"{prefix}{separator}{next_number}"
        while replacement in reserved:
            next_number += 1
            replacement = f"{prefix}{separator}{next_number}"
        shape.set("id", replacement)
        if shape.get(office_spid) == shape_id:
            shape.set(office_spid, replacement)
        seen.add(replacement)
        reserved.add(replacement)
        changed = True
    if not changed:
        return content
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)


def chinese_date(value):
    if not value:
        return "____年__月__日"
    year, month, day = value.split("-")
    return f"{year}年{int(month)}月{int(day)}日"


def generate(snapshot, root, *, draft=False):
    test = snapshot["test"]
    method = test["method"]
    selected_template = template_name(test, snapshot["attachments"])
    document = Document(TEMPLATES / f"{selected_template}.docx")
    manifest = json.loads((TEMPLATES / "manifest.json").read_text(encoding="utf-8"))[selected_template]
    body = list(document.element.body)
    for element in body:
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            if any(label in paragraph.text for label in ("样本统计表", "原始记录", "显微观察记录", "实验结果统计表")):
                paragraph.paragraph_format.keep_with_next = True
    tables = [el for el in body if el.tag == qn("w:tbl")]
    metadata = Table(tables[0], document)
    cell_text(
        metadata.cell(1, 0),
        f"1、采样日期：{chinese_date(test['samplingDate'])}         实验日期：{chinese_date(test['testDate'])}",
    )
    materials = test.get("reportMaterial") or "、".join(dict.fromkeys(s["material"] for s in test["samples"]))
    states = test.get("reportSpecimenState") or "、".join(
        dict.fromkeys(s.get("specimenState", "") for s in test["samples"] if s.get("specimenState"))
    )
    pools = sum(s["poolCount"] for s in test["samples"])
    portions = sum(s["portionCount"] for s in test["samples"])
    quantity_material = materials if test.get("reportFormVersion") == 2 else ""
    cell_text(
        metadata.cell(2, 0),
        f"样品名称：{materials}     样品状态：{states}     样品数量：{pools}样（{portions}份{quantity_material}）",
    )
    sample_element = next(t for t in tables[1:] if Table(t, document).cell(0, 0).text == "样本编号")
    source_map = {s["id"]: s for s in snapshot["batch"]["sources"]}
    replace_element(
        sample_element, [sample_table(sample_element, document, test["samples"], source_map, pcr=method == "pcr")]
    )
    if method.startswith("elisa"):
        kit_element = tables[1]
        replace_element(kit_element, [reagent_table(kit_element, document, test["projects"])])
    elif method == "pcr":
        for element in body:
            if element.tag != qn("w:p"):
                continue
            paragraph = Paragraph(element, document)
            if "{{kit}}" in paragraph.text:
                paragraph_text(
                    paragraph, "试剂盒名称：" + "；".join(dict.fromkeys(p.get("kit", "") for p in test["projects"]))
                )
            if "{{lot}}" in paragraph.text:
                paragraph_text(
                    paragraph, "试剂盒批号：" + "；".join(dict.fromkeys(p.get("lot", "") for p in test["projects"]))
                )
    entries = image_entries(test, snapshot["attachments"])
    result_heading = next(el for el in body if el.tag == qn("w:p") and "实验结果统计表" in Paragraph(el, document).text)
    pictures = [
        el
        for el in body[: body.index(result_heading)]
        if el.tag == qn("w:tbl") and not Table(el, document).cell(0, 0).text
    ]
    offset = 0
    for index, picture in enumerate(pictures):
        table = Table(picture, document)
        capacity = len(table.rows) // 2 * len(table.columns)
        chunk = entries[offset:] if index == len(pictures) - 1 else entries[offset : offset + capacity]
        replacement = (
            [
                image_table(
                    picture, document, chunk, root, manifest["imageSizes"][offset:] or manifest["imageSizes"][-1:]
                )
            ]
            if chunk or index == 0
            else []
        )
        replace_element(picture, replacement)
        offset += capacity
    legend = next(el for el in body if el.tag == qn("w:p") and Paragraph(el, document).text.startswith("注："))
    region = body[body.index(result_heading) + 1 : body.index(legend)]
    patterns = result_patterns(region, document)
    prefix = result_prefix(patterns, document)
    for element in region:
        element.getparent().remove(element)
    for index, project in enumerate(test["projects"]):
        heading_pattern, pattern = sized_result_pattern(patterns, document, project, controls=method != "parasite")
        heading = copy_paragraph(heading_pattern, document, f"{prefix}-{index + 1}{project['name']}")
        Paragraph(heading, document).paragraph_format.keep_with_next = True
        legend.addprevious(heading)
        for element in horizontal_result(pattern, document, project, test["samples"], controls=method != "parasite"):
            legend.addprevious(element)
    signature = next(
        el
        for el in body
        if el.tag == qn("w:p") and Paragraph(el, document).text.strip().startswith(("检测人：", "检验人："))
    )
    # Empty filler paragraphs in the examples located the signature; dynamic content flows naturally instead.
    for element in body[body.index(legend) + 1 : body.index(signature)]:
        if element.tag == qn("w:p") and not Paragraph(element, document).text.strip():
            element.getparent().remove(element)
    if test.get("reportFormVersion") != 2 and test["conclusion"]:
        conclusion = copy_paragraph(legend, document, "结果判定：" + test["conclusion"])
        signature.addprevious(conclusion)
    if test.get("notes"):
        signature.addprevious(copy_paragraph(legend, document, "备注：" + test["notes"]))
    for attachment in snapshot["attachments"]:
        if not attachment["mime"].startswith("image/"):
            signature.addprevious(copy_paragraph(legend, document, "原始附件（系统留档）：" + attachment["name"]))
    # Keep the short closing block with its signature instead of orphaning the signature on a page.
    Paragraph(legend, document).paragraph_format.keep_with_next = True
    # Bring the final result table along when the closing block needs a new page.
    previous = legend.getprevious()
    while previous is not None:
        if previous.tag == qn("w:p"):
            Paragraph(previous, document).paragraph_format.keep_with_next = True
        elif previous.tag == qn("w:tbl"):
            for row in Table(previous, document).rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        paragraph.paragraph_format.keep_with_next = True
            break
        previous = previous.getprevious()
    closing = signature.getprevious()
    while closing is not None and closing is not legend:
        if closing.tag == qn("w:p"):
            Paragraph(closing, document).paragraph_format.keep_with_next = True
        closing = closing.getprevious()
    if draft:
        for section in document.sections:
            for footer in (section.footer, section.first_page_footer, section.even_page_footer):
                footer.paragraphs[0].add_run("  草稿 · 仅供预览")
    # Detached table clones are invisible to python-docx's next-ID allocator.
    for index, drawing in enumerate(document.element.xpath(".//wp:docPr"), 1):
        drawing.set("id", str(index))
    output = BytesIO()
    document.save(output)
    return preserve_parts(output.getvalue(), TEMPLATES / f"{selected_template}.docx", draft=draft)


def preserve_parts(content, template, *, draft):
    """Keep untouched source parts byte-for-byte, including their XML serialization."""
    output = BytesIO()
    with ZipFile(template) as original, ZipFile(BytesIO(content)) as generated, ZipFile(output, "w") as result:
        retained = {
            name
            for name in original.namelist()
            if name in {"word/styles.xml", "word/numbering.xml", "word/fontTable.xml"}
            or name.startswith(("word/theme/", "word/header", "word/_rels/header"))
            or (not draft and name.startswith(("word/footer", "word/_rels/footer")))
        }
        for entry in generated.infolist():
            source = original if entry.filename in retained else generated
            part = source.read(entry.filename)
            if entry.filename == "word/styles.xml":
                part = normalize_style_order(part)
            elif entry.filename.startswith("word/footer") and entry.filename.endswith(".xml"):
                part = normalize_footer_shape_ids(part)
            result.writestr(entry, part)
    return output.getvalue()


def image_entries(test, attachments):
    images = sorted(
        [a for a in attachments if a["mime"].startswith("image/") and not a.get("removed")],
        key=lambda a: (a.get("position", 0), a.get("uploadedAt", a.get("updatedAt", "")), a["id"]),
    )
    entries, used = [], set()
    if test["method"] == "parasite":
        for sample in test["samples"]:
            grouped = [
                [a for a in images if a.get("sampleId") == sample["id"] and a.get("category") == category]
                for category in ("体外", "体内")
            ]
            for index in range(max(1, *(len(group) for group in grouped))):
                for category, group in zip(("体外", "体内"), grouped, strict=True):
                    attachment = group[index] if index < len(group) else None
                    if attachment:
                        used.add(attachment["id"])
                    entries.append(
                        {
                            "attachment": attachment,
                            "caption": " / ".join(
                                filter(
                                    None,
                                    [
                                        f"{sample['number']}{category}寄生虫",
                                        attachment.get("caption") if attachment else "",
                                    ],
                                )
                            ),
                        }
                    )
    for attachment in images:
        if attachment["id"] in used:
            continue
        project = "、".join(p["name"] for p in test["projects"] if p["id"] in project_ids(attachment))
        sample = next((s["number"] for s in test["samples"] if s["id"] == attachment.get("sampleId")), "")
        entries.append(
            {
                "attachment": attachment,
                "caption": attachment.get("caption")
                or " / ".join(filter(None, [project or attachment["name"], sample])),
            }
        )
    return entries
