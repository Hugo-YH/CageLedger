"""Render a batch overview using the retained six-column laboratory report."""

from collections import Counter, defaultdict
from copy import deepcopy
from io import BytesIO

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.table import _Row

from server_app.domains.intake.rules import species_label

from . import repository as repo
from .catalog import PROJECTS
from .document_layout import cell_text, compact_cell, paragraph_text
from .documents import TEMPLATES, chinese_date, preserve_parts

TITLES = {"quarantine": "实验动物检疫检测报告", "self": "实验动物自检检测报告"}
# Only unambiguous historical names match a catalog project. MVM is not MPV;
# a general internal-parasite result cannot establish individual protozoa results.
ROW_PROJECTS = {
    2: ("elisa_mouse", "小鼠仙台病毒抗体（SV）"),
    3: ("elisa_mouse", "小鼠肝炎病毒抗体（MHV）"),
    4: ("elisa_mouse", "小鼠肺炎病毒抗体（PVM）"),
    5: ("elisa_mouse", "小鼠呼肠孤病毒Ⅲ型抗体（Reo-3）"),
    9: ("elisa_rat", "大鼠仙台病毒抗体（SV）"),
    10: ("elisa_rat", "大鼠呼肠孤病毒Ⅲ抗体（Reo-3）"),
    11: ("elisa_rat", "大鼠肺炎病毒抗体（PVM）"),
    12: ("elisa_rat", "大鼠汉坦病毒抗体（HV）"),
    15: ("pcr", "沙门菌"),
    16: ("elisa_mouse", "小鼠肺支原体抗体（MYco）"),
    17: ("pcr", "鼠棒状杆菌"),
    18: ("elisa_mouse", "小鼠泰泽病原体抗体"),
    20: ("pcr", "肺炎克雷伯杆菌"),
    21: ("pcr", "绿脓杆菌"),
    30: ("parasite", "体外寄生虫"),
    34: ("elisa_mouse", "小鼠弓形虫抗体（Toxo）"),
}


def batch_snapshot(conn, batch_id):
    batch = repo.get(conn, "batches", batch_id)
    tests = repo.all_items(conn, "tests", batch_id)
    superseded = {test.get("correctionOf") for test in tests}
    records = []
    for test in sorted(tests, key=lambda t: (t["testDate"], t["method"], t["id"])):
        if test["id"] in superseded:
            continue
        reports = repo.all_items(conn, "reports", test["id"])
        if test["state"] == "issued" and reports:
            report = max(reports, key=lambda r: r["version"])
            records.append({"test": report["snapshot"]["test"], "number": report["number"], "issued": True})
        else:
            records.append({"test": test, "number": "", "issued": False})
    if not records:
        raise ValueError("请先保存检测记录，再导出汇总报告")
    return {"batch": batch, "records": records}


def project_result(test, project):
    applicable = set(project["sampleIds"]) & {sample["id"] for sample in test["samples"]}
    counts = Counter(project["results"].get(sid, "") for sid in applicable)
    tested = counts["negative"] + counts["positive"] + counts["suspect"]
    result = f"{counts['positive']}/{tested}" if tested else "—"
    notes = [
        f"{count}组{label}"
        for key, label in [("suspect", "可疑"), ("not_tested", "未检测"), ("", "未填写")]
        if (count := counts[key])
    ]
    if not applicable:
        notes.append("无适用实验组")
    if test.get("retestOf"):
        notes.append("复检")
    if test.get("correctionOf"):
        notes.append("更正")
    return result, "；".join(notes)


def category(method, name):
    if name not in PROJECTS.get(method, []):
        return 3
    if method == "parasite" or "弓形虫" in name:
        return 2
    if method == "pcr" or any(word in name for word in ("支原体", "泰泽")):
        return 1
    return 0


def generate(snapshot, *, kind="quarantine"):
    if kind not in TITLES:
        raise ValueError("报告类型无效")
    template = TEMPLATES / "summary.docx"
    document = Document(template)
    records = snapshot["records"]
    draft = any(not record["issued"] for record in records)
    tests = [record["test"] for record in records]
    species = "、".join(
        dict.fromkeys(species_label(s.get("species", "")) for s in snapshot["batch"]["sources"] if s.get("species"))
    )
    dates = "、".join(chinese_date(date) for date in sorted({t["samplingDate"] for t in tests}))
    for paragraph in document.paragraphs:
        if paragraph.text == "实验动物检疫检测报告":
            paragraph_text(paragraph, TITLES[kind])
        elif paragraph.text == "{{species_date}}":
            paragraph_text(paragraph, f"动物种类：{species or '未填写'}    采样时间：{dates}")
        elif paragraph.text.startswith("“-”代表"):
            paragraph_text(paragraph, "“—”表示本次无已检测结果；项目要求标记沿用历史表，不表示本次已完成检测。")
    entries = defaultdict(list)
    for test in tests:
        for project in test["projects"]:
            entries[(test["method"], project["name"])].append((test, project))
    table = document.tables[0]
    # Keep dates and method names on one line, with room for incomplete/retest remarks.
    widths = [column.width for column in table.columns]
    adjustments = [-720000, 0, 0, 180000, 180000, 360000]
    for column, width, adjustment in zip(table.columns, widths, adjustments, strict=True):
        column.width = width + adjustment
    for row in table.rows:
        if len(row._tr.tc_lst) > 1:
            for cell, width, adjustment in zip(row.cells, widths, adjustments, strict=True):
                cell.width = width + adjustment
    table.autofit = False
    result_dates = {test["testDate"] for test in tests}
    single_date = next(iter(result_dates)) if len(result_dates) == 1 else ""
    if single_date:
        cell_text(table.cell(0, 4), single_date.replace("-", "/") + "\n检测结果")
    rows = list(table.rows)
    for index, key in ROW_PROJECTS.items():
        if key in entries:
            fill_row(rows[index], entries.pop(key), single_date=single_date)
    # Preserve all original rows; add precise projects rather than merge unlike assays.
    other_heading = None
    if any(category(method, name) == 3 for method, name in entries):
        other_heading = deepcopy(rows[29]._tr)
        table._tbl.append(other_heading)
        cell_text(_Row(other_heading, table).cells[0], "其他检测项目")
    anchors = [rows[14]._tr, rows[29]._tr, other_heading, None]
    for (method, name), values in entries.items():
        element = deepcopy(rows[2]._tr)
        group = category(method, name)
        if anchors[group] is None:
            table._tbl.append(element)
        else:
            anchors[group].addprevious(element)
        row = _Row(element, table)
        cell_text(row.cells[0], name)
        cell_text(row.cells[1], "")
        fill_row(row, values, single_date=single_date)
    for row in table.rows:
        if row.height_rule == WD_ROW_HEIGHT_RULE.EXACTLY:
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            compact_cell(cell)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    signature = next(p for p in document.paragraphs if p.text.startswith("检测人："))
    status = "汇总草稿 · 含未出具记录，仅供核对" if draft else "汇总导出 · 根据已出具报告汇编，签名栏由人员签署"
    provenance = "；".join(record["number"] or f"{record['test']['testDate']}草稿" for record in records)
    paragraph_text(signature.insert_paragraph_before(), f"{status}。来源：{provenance}。")
    if draft:
        for section in document.sections:
            section.footer.paragraphs[0].add_run("  草稿 · 仅供预览")
    output = BytesIO()
    document.save(output)
    return preserve_parts(output.getvalue(), template, draft=draft)


def fill_row(row, entries, *, single_date=""):
    results, notes, specimens, methods = [], [], [], []
    for test, project in entries:
        value, note = project_result(test, project)
        date = test["testDate"] or "日期未填"
        if test.get("retestOf"):
            value = "复检 " + value
        elif test.get("correctionOf"):
            value = "更正 " + value
        results.append(value if single_date else f"{date}\n{value}")
        if note:
            notes.append(note if single_date else f"{date}：{note}")
        applicable = set(project["sampleIds"])
        specimens.extend(s["material"] for s in test["samples"] if s["id"] in applicable)
        methods.append("ELISA" if test["method"].startswith("elisa") else "PCR" if test["method"] == "pcr" else "Micro")
    cell_text(row.cells[2], "、".join(dict.fromkeys(specimens)) or "—")
    cell_text(row.cells[3], "、".join(dict.fromkeys(methods)))
    cell_text(row.cells[4], "\n".join(results))
    cell_text(row.cells[5], "\n".join(notes))
