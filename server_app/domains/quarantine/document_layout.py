"""Fill cloned source-template elements without replacing their formatting."""

from copy import deepcopy
from io import BytesIO

from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

SYMBOLS = {"negative": "-", "positive": "＋", "suspect": "±", "not_tested": "未检测", "": "/"}


def paragraph_text(paragraph, value):
    runs = paragraph.runs
    if runs:
        runs[0].text = str(value)
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(str(value))


def cell_text(cell, value):
    paragraph_text(cell.paragraphs[0], value)
    for paragraph in cell.paragraphs[1:]:
        cell._tc.remove(paragraph._p)


def copy_paragraph(element, doc, text):
    clone = deepcopy(element)
    paragraph_text(Paragraph(clone, doc), text)
    return clone


def replace_element(old, elements):
    for element in elements:
        old.addprevious(element)
    old.getparent().remove(old)


def sample_table(template, doc, samples, sources, *, pcr=False):
    result = deepcopy(template)
    table = Table(result, doc)
    group_size = 3 if pcr else 2
    patterns = _row_groups(table, group_size)
    for row in list(table.rows):
        result.remove(row._tr)
    offset, group_index = 0, 0
    while offset < max(len(samples), 1):
        pattern = patterns[min(group_index, len(patterns) - 1)]
        for row in pattern:
            result.append(deepcopy(row))
        rows = Table(result, doc).rows[-group_size:]
        # Some historical PCR forms omit the last cell in the second row group.
        capacity = min(len(row.cells) for row in rows) - 1
        for index in range(capacity):
            sample = samples[offset + index] if offset + index < len(samples) else None
            cell_text(rows[0].cells[index + 1], sample["number"] if sample else "")
            selected = [sources[sid] for sid in sample["sourceIds"]] if sample else []
            supplier_names = list(dict.fromkeys(s["supplier"] for s in selected))
            people = [
                "".join(filter(None, [s["pi"], s["owner"]]))
                + ("\n" + s.get("notes", "") if s.get("manual") and s.get("notes") else "")
                for s in selected
            ]
            if pcr:
                cell_text(rows[1].cells[index + 1], "\n".join(supplier_names))
                cell_text(rows[2].cells[index + 1], "\n".join(filter(None, people)))
            else:
                cell_text(rows[1].cells[index + 1], "\n".join([*supplier_names, *filter(None, people)]))
        offset += capacity
        group_index += 1
    for row_index, row in enumerate(Table(result, doc).rows):
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if row_index % group_size < group_size - 1:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
    unfloat(result)
    return result


def horizontal_result(template, doc, project, samples, *, controls):
    source = Table(template, doc)
    capacity = max(1, len(source.columns) - 1 - (2 if controls else 0))
    applicable = [s for s in samples if s["id"] in project["sampleIds"]]
    elements = []
    for offset in range(0, max(len(applicable), 1), capacity):
        chunk = applicable[offset : offset + capacity]
        columns = [(s["id"], s["number"]) for s in chunk]
        if controls:
            columns.extend([("nc", "NC"), ("pc", "PC")])
        if not columns:
            columns = [("", "无适用样本")]
        result = deepcopy(template)
        table = Table(result, doc)
        grid = table._tbl.tblGrid
        original_grid = list(grid)
        for col in list(grid):
            grid.remove(col)
        grid.append(deepcopy(original_grid[0]))
        source_indices = [
            (-2 if sid == "nc" else -1 if sid == "pc" else min(index + 1, len(original_grid) - 1))
            for index, (sid, _) in enumerate(columns)
        ]
        for source_index in source_indices:
            grid.append(deepcopy(original_grid[source_index]))
        for row in table.rows:
            cells = list(row._tr.tc_lst)
            for c in cells[1:]:
                row._tr.remove(c)
            for source_index in source_indices:
                row._tr.append(deepcopy(cells[source_index]))
        table = Table(result, doc)
        for index, (sid, label) in enumerate(columns, 1):
            cell_text(table.rows[0].cells[index], label)
            if len(table.rows) == 3:
                cell_text(table.rows[1].cells[index], project.get("wells", {}).get(sid, ""))
            value = project.get(sid, "") if sid in {"nc", "pc"} else project["results"].get(sid, "")
            cell_text(table.rows[-1].cells[index], SYMBOLS.get(value, "未填写") if sid else "/")
        repeat = OxmlElement("w:tblHeader")
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in table.rows[:-1]:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.keep_with_next = True
        unfloat(result)
        elements.append(result)
        if offset + capacity < len(applicable):
            # Only consecutive chunks need a separator; the next project has its own heading.
            # An empty normal-sized paragraph after every project adds unwanted vertical gaps.
            separator = OxmlElement("w:p")
            paragraph = Paragraph(separator, doc)
            paragraph.paragraph_format.space_after = 0
            paragraph.paragraph_format.space_before = 0
            paragraph.paragraph_format.line_spacing = Pt(1)
            elements.append(separator)
    return elements


def reagent_table(template, doc, projects):
    result = deepcopy(template)
    table = Table(result, doc)
    pattern = deepcopy(table.rows[1]._tr)
    for row in list(table.rows)[1:]:
        result.remove(row._tr)
    seen = set()
    for project in projects:
        key = (project.get("kit", ""), project.get("lot", ""))
        if key in seen:
            continue
        seen.add(key)
        result.append(deepcopy(pattern))
        row = Table(result, doc).rows[-1]
        cell_text(row.cells[0], key[0])
        cell_text(row.cells[1], key[1])
    unfloat(result)
    return result


def image_table(template, doc, entries, root, sizes):
    result = deepcopy(template)
    source = Table(result, doc)
    source.autofit = False
    capacity = len(source.columns)
    patterns = _row_groups(source, 2)
    for row in list(source.rows):
        result.remove(row._tr)
    for group_index, offset in enumerate(range(0, max(len(entries), capacity), capacity)):
        pattern = patterns[min(group_index, len(patterns) - 1)]
        for row in pattern:
            result.append(deepcopy(row))
        rows = Table(result, doc).rows[-2:]
        for row in rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for index in range(capacity):
            entry = entries[offset + index] if offset + index < len(entries) else None
            caption = entry["caption"] if entry else ""
            cell_text(rows[1].cells[index], caption)
            cell = rows[0].cells[index]
            cell_text(cell, "")
            if entry and entry.get("attachment"):
                attachment = entry["attachment"]
                shape = (
                    cell.paragraphs[0]
                    .add_run()
                    .add_picture(
                        BytesIO((root / attachment.get("previewStorageName", attachment["storageName"])).read_bytes())
                    )
                )
                image_width, height = _image_size(sizes, offset + index)
                width = cell.width or image_width
                ratio = min(min(image_width, max(width - 150000, 1)) / shape.width, height / shape.height)
                shape.width, shape.height = int(shape.width * ratio), int(shape.height * ratio)
            elif entry:
                cell_text(cell, "未上传图片")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    unfloat(result)
    return result


def _row_groups(table, group_size):
    rows = [deepcopy(row._tr) for row in table.rows]
    groups = [rows[index : index + group_size] for index in range(0, len(rows), group_size)]
    return [group for group in groups if len(group) == group_size]


def _image_size(sizes, index):
    if not sizes:
        return 2700000, 2100000
    return sizes[min(index, len(sizes) - 1)]


def unfloat(table):
    for floating in table.xpath("./w:tblPr/w:tblpPr"):
        floating.getparent().remove(floating)
