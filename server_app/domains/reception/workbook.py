from copy import copy
from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TEMPLATE = Path(__file__).resolve().parents[2] / "resources" / "intake" / "reception-register.xlsx"
STYLE_ATTRIBUTES = ("font", "fill", "border", "alignment", "protection", "number_format")
MISSING_VALUE = "待补充"
MISSING_FILL = "FFFFC7CE"
MISSING_TEXT = "FF9C0006"
HEADERS = (
    "采购订单编号",
    "项目负责人",
    "IACUC编号",
    "到期时间",
    "饲养需求批次号",
    "项目来源",
    "接收时间",
    "实验负责人",
    "具体品系",
    "合格证品系",
    "等级",
    "",
    "体重日龄",
    "性别",
    "数量（只）",
    "公司",
    "动物质量合格证号",
    "动物实验证明编号",
    "饲养房间",
    "备注",
)


def generate(document):
    workbook = load_workbook(TEMPLATE)
    sheet = workbook.active
    default_height = sheet.sheet_format.defaultRowHeight
    body_height = sheet.row_dimensions[4].height or default_height
    sheet.row_dimensions[3].height = sheet.row_dimensions[3].height or default_height
    row_styles = [
        {attribute: copy(getattr(sheet.cell(4, column), attribute)) for attribute in STYLE_ATTRIBUTES}
        for column in range(2, 22)
    ]
    for column, label in enumerate(HEADERS, 2):
        sheet.cell(3, column, label)
    for index, line in enumerate(document["rows"], 4):
        sheet.row_dimensions[index].height = body_height
        project = line["project"]
        values = [
            line["purchaseOrderNo"],
            project["pi"],
            project["iacuc"],
            project["projectEndDate"],
            line["batchNo"],
            project["funding"],
            line["actualReceiptDate"],
            project["owner"],
            line["strain"],
            line["certificateStrain"],
            line["grade"],
            "",
            line["specification"],
            line["sex"],
            line["quantity"],
            line["supplier"],
            line["qualityCertificateNo"],
            line["experimentCertificateNo"],
            line["roomName"],
            line["notes"],
        ]
        for column, value in enumerate(values, 2):
            cell = sheet.cell(index, column)
            original = row_styles[column - 2]
            for attribute in STYLE_ATTRIBUTES:
                setattr(cell, attribute, copy(original[attribute]))
            if column in (5, 8) and value:
                try:
                    cell.value = date.fromisoformat(str(value))
                    continue
                except ValueError:
                    value = ""
            if column not in (13, 21) and (value is None or isinstance(value, str) and not value.strip()):
                value = MISSING_VALUE
                cell.fill = PatternFill("solid", fgColor=MISSING_FILL)
                font = copy(cell.font)
                font.color = MISSING_TEXT
                cell.font = font
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"
    sheet.auto_filter.ref = f"B3:U{max(3, sheet.max_row)}"
    notes = workbook.create_sheet("导出说明")
    for value in (
        "接收登记预填表：请核对实际只数、性别及证件；复制B4起的数据到原登记表B列。",
        "项目字段以实验申请汇总表为准，匹配失败不使用预约信息兜底。",
        "数量、性别及规格是预约预填；合格证信息尚未线上登记，请核对后补全。",
        "等级按物种预填：小鼠、大鼠为SPF，其他动物为普通级；物种未填写时等级标记待补充。",
        "缺失字段以红底“待补充”标记；空白M列和可空备注保持空白，数值0不视为缺失。",
        "不输出V列以后的入库报销数据，也不推进批次状态。",
        f"申请资料版本：{document['applicationVersion']}",
        *document["warnings"],
    ):
        cell = notes.cell(notes.max_row + (1 if notes.cell(1, 1).value else 0), 1, value)
        cell.data_type = "s"
        cell.font = Font(name="宋体", size=20)
        cell.alignment = Alignment(wrap_text=True)
    notes.column_dimensions["A"].width = 100
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()
