import copy
import tempfile
import unittest
from collections import Counter
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from lxml import etree

from server_app.domains.quarantine.catalog import PROJECTS
from server_app.domains.quarantine.document_patterns import template_name
from server_app.domains.quarantine.documents import (
    OFFICE,
    TEMPLATES,
    VML,
    WORDPROCESSINGML,
    generate,
    image_entries,
    normalize_footer_shape_ids,
    normalize_style_order,
)


def report_snapshot(method, count):
    samples = [
        {
            "id": f"sample-{index}",
            "number": str(index + 1),
            "sourceIds": ["source"],
            "material": "血清" if method.startswith("elisa") else "粪便",
            "poolCount": 1,
            "portionCount": 2,
        }
        for index in range(count)
    ]
    return {
        "test": {
            "method": method,
            "samplingDate": "2026-09-02",
            "testDate": "2026-09-03",
            "reportFormVersion": 2,
            "reportSpecimenState": "合格",
            "conclusion": "",
            "samples": samples,
            "projects": [
                {
                    "id": f"project-{index}",
                    "name": name,
                    "sampleIds": [sample["id"] for sample in samples],
                    "results": {sample["id"]: "negative" for sample in samples},
                    "wells": {sample["id"]: str(i + 1) for i, sample in enumerate(samples)},
                    "nc": "negative",
                    "pc": "positive",
                    "kit": f"检测用试剂盒{index + 1}",
                    "lot": "QA-202609",
                }
                for index, name in enumerate(PROJECTS[method])
            ],
        },
        "batch": {"sources": [{"id": "source", "supplier": "测试供应商", "pi": "甲", "owner": "乙"}]},
        "attachments": [],
    }


def result_table(document, project):
    heading = next(p for p in document.paragraphs if p.text.endswith(project))
    element = heading._p.getnext()
    while element.tag != qn("w:tbl"):
        element = element.getnext()
    return Table(element, document)


def style_children(content):
    root = etree.fromstring(content)
    style_tag = f"{{{WORDPROCESSINGML}}}style"
    return [[etree.tostring(child) for child in style] for style in root.findall(style_tag)]


def assert_valid_style_order(test_case, content):
    root = etree.fromstring(content)
    style_tag = f"{{{WORDPROCESSINGML}}}style"
    ui_priority_tag = f"{{{WORDPROCESSINGML}}}uiPriority"
    after_ui_priority = {
        f"{{{WORDPROCESSINGML}}}semiHidden",
        f"{{{WORDPROCESSINGML}}}unhideWhenUsed",
        f"{{{WORDPROCESSINGML}}}qFormat",
    }
    for style in root.findall(style_tag):
        ui_priority = style.find(ui_priority_tag)
        if ui_priority is None:
            continue
        test_case.assertLess(
            style.index(ui_priority),
            min((style.index(child) for child in style if child.tag in after_ui_priority), default=len(style)),
            style.get(f"{{{WORDPROCESSINGML}}}styleId"),
        )


def footer_shape_ids(content):
    root = etree.fromstring(content)
    return [shape.get("id") for shape in root.iter(f"{{{VML}}}shape") if shape.get("id")]


def footer_shape_spids(content):
    root = etree.fromstring(content)
    return [shape.get(f"{{{OFFICE}}}spid") for shape in root.iter(f"{{{VML}}}shape")]


class QuarantineDocumentTemplateTests(unittest.TestCase):
    def test_explicit_image_caption_is_preserved_without_repeating_project_names(self):
        snapshot = report_snapshot("pcr", 1)
        attachment = {
            "id": "image",
            "mime": "image/png",
            "name": "image.png",
            "projectIds": ["project-0"],
            "sampleId": "sample-0",
            "caption": "绿脓杆菌",
        }
        self.assertEqual(image_entries(snapshot["test"], [attachment])[0]["caption"], "绿脓杆菌")
        attachment["caption"] = ""
        self.assertEqual(image_entries(snapshot["test"], [attachment])[0]["caption"], "绿脓杆菌 / 1")

    def test_all_methods_retain_heading_legend_signature_and_page_furniture(self):
        cases = [
            ("parasite", 5, "parasite"),
            ("parasite", 6, "parasite_6"),
            ("parasite", 7, "parasite_7"),
            ("elisa_mouse", 3, "elisa_mouse"),
            ("elisa_mouse", 4, "elisa_mouse_4"),
            ("elisa_mouse", 7, "elisa_mouse_7"),
            ("elisa_rat", 1, "elisa_rat_1"),
            ("elisa_rat", 2, "elisa_rat"),
            ("pcr", 6, "pcr_small"),
            ("pcr", 8, "pcr"),
            ("pcr", 9, "pcr_9"),
            ("pcr", 10, "pcr_9"),
        ]
        for method, count, selected in cases:
            with self.subTest(method=method, count=count), tempfile.TemporaryDirectory() as temporary:
                snapshot = report_snapshot(method, count)
                self.assertEqual(template_name(snapshot["test"], []), selected)
                original = Document(TEMPLATES / f"{selected}.docx")
                output = generate(snapshot, Path(temporary))
                exported = Document(BytesIO(output))
                self.assertEqual(exported.sections[0]._sectPr.xml, original.sections[0]._sectPr.xml)
                for marker in ("实验结果统计表", "注：", "检测人："):
                    baseline = next(p for p in original.paragraphs if marker in p.text)
                    actual = next(p for p in exported.paragraphs if marker in p.text)
                    self.assertEqual(actual.text, baseline.text)
                    self.assertEqual(
                        [e.xml for e in actual._p.xpath("./w:pPr/w:numPr")],
                        [e.xml for e in baseline._p.xpath("./w:pPr/w:numPr")],
                    )
                with ZipFile(BytesIO(output)) as generated, ZipFile(TEMPLATES / f"{selected}.docx") as source:
                    for name in source.namelist():
                        if name == "word/styles.xml":
                            self.assertEqual(generated.read(name), normalize_style_order(source.read(name)), name)
                        elif name.startswith("word/footer") and name.endswith(".xml"):
                            self.assertEqual(generated.read(name), normalize_footer_shape_ids(source.read(name)), name)
                        elif name == "word/numbering.xml" or name.startswith("word/header"):
                            self.assertEqual(generated.read(name), source.read(name), name)
                    self.assertNotIn("{{", generated.read("word/document.xml").decode())
                prefix = "5" if method == "pcr" else "4"
                self.assertTrue(any(p.text.startswith(f"{prefix}-1") for p in exported.paragraphs))

    def test_reordered_projects_use_their_own_result_tables(self):
        snapshot = report_snapshot("pcr", 8)
        # The historical rat-only table has a different width from the general table.
        rat = copy.deepcopy(snapshot["test"]["projects"][-1])
        rat["sampleIds"] = rat["sampleIds"][:3]
        snapshot["test"]["projects"] = [rat, snapshot["test"]["projects"][0]]
        with tempfile.TemporaryDirectory() as temporary:
            exported = Document(BytesIO(generate(snapshot, Path(temporary))))
        original = Document(TEMPLATES / "pcr.docx")
        for project in snapshot["test"]["projects"]:
            expected = result_table(original, project["name"])
            actual = result_table(exported, project["name"])
            self.assertEqual(actual.cell(0, 0).width, expected.cell(0, 0).width)
            self.assertEqual(actual.cell(0, 1).width, expected.cell(0, 1).width)
            self.assertEqual(actual.rows[0].cells[-2].text, "NC")
            self.assertEqual(actual.rows[0].cells[-1].text, "PC")
        self.assertEqual(len(result_table(exported, rat["name"]).columns), 6)

    def test_new_project_uses_full_width_pattern_and_preserves_all_samples(self):
        snapshot = report_snapshot("pcr", 8)
        project = snapshot["test"]["projects"][0]
        project["name"] = "自定义检测项目"
        snapshot["test"]["projects"] = [project]
        with tempfile.TemporaryDirectory() as temporary:
            exported = Document(BytesIO(generate(snapshot, Path(temporary))))
        table = result_table(exported, project["name"])
        self.assertEqual([cell.text for cell in table.rows[0].cells[1:]], [*map(str, range(1, 9)), "NC", "PC"])
        self.assertEqual([cell.text for cell in table.rows[-1].cells[1:]], [*(["-"] * 9), "＋"])

    def test_large_mouse_cohort_uses_one_result_table_per_project(self):
        snapshot = report_snapshot("elisa_mouse", 7)
        with tempfile.TemporaryDirectory() as temporary:
            exported = Document(BytesIO(generate(snapshot, Path(temporary))))
        self.assertEqual(len(exported.tables), 4 + len(PROJECTS["elisa_mouse"]))
        sample_table = exported.tables[2]
        self.assertEqual(len(sample_table.rows), 2)
        self.assertEqual([c.text for c in sample_table.rows[0].cells[1:]], list(map(str, range(1, 8))))
        for project in snapshot["test"]["projects"]:
            actual = result_table(exported, project["name"])
            self.assertEqual([c.text for c in actual.rows[0].cells[1:]], [*map(str, range(1, 8)), "NC", "PC"])

    def test_one_elisa_image_selects_the_retained_single_column_layout(self):
        from tests.test_quarantine_document_layout import PNG

        for count, expected in [(4, "elisa_mouse_4_single"), (8, "elisa_mouse_8_single")]:
            snapshot = report_snapshot("elisa_mouse", count)
            self.assertEqual(template_name(snapshot["test"], [{"mime": "image/png"}]), expected)
            self.assertNotEqual(template_name(snapshot["test"], [{"mime": "image/png", "removed": True}]), expected)
            snapshot["attachments"] = [{"id": "image", "mime": "image/png", "name": "QA", "storageName": "qa.png"}]
            with tempfile.TemporaryDirectory() as temporary:
                root = Path(temporary)
                (root / "qa.png").write_bytes(PNG)
                exported = Document(BytesIO(generate(snapshot, root)))
            self.assertEqual(len(exported.inline_shapes), 1)
            self.assertEqual(len(exported.tables[3].columns), 1)
            signature = next(p.text for p in exported.paragraphs if "复核人：" in p.text)
            original = Document(TEMPLATES / f"{expected}.docx")
            self.assertEqual(signature, next(p.text for p in original.paragraphs if "复核人：" in p.text))

    def test_pcr_omitted_grid_cell_and_overflow_preserve_every_sample(self):
        snapshot = report_snapshot("pcr", 12)
        with tempfile.TemporaryDirectory() as temporary:
            exported = Document(BytesIO(generate(snapshot, Path(temporary))))
        rows = exported.tables[1].rows
        self.assertEqual([len(row.cells) for row in rows[::3]], [6, 5, 5])
        self.assertEqual([c.text for row in rows[::3] for c in row.cells[1:] if c.text], list(map(str, range(1, 13))))

    def test_multi_table_images_have_unique_drawing_ids_and_preserve_captions(self):
        from tests.test_quarantine_document_layout import PNG

        snapshot = report_snapshot("pcr", 9)
        snapshot["attachments"] = [
            {"id": str(i), "mime": "image/png", "name": f"image-{i}", "storageName": "qa.png", "position": i}
            for i in range(6)
        ]
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "qa.png").write_bytes(PNG)
            exported = Document(BytesIO(generate(snapshot, root)))
        identifiers = exported.element.xpath(".//wp:docPr/@id")
        self.assertEqual(len(identifiers), 6)
        self.assertEqual(len(set(identifiers)), 6)
        captions = [c.text for table in exported.tables[2:4] for row in table.rows[1::2] for c in row.cells]
        self.assertEqual(captions, [f"image-{i}" for i in range(6)])

    def test_added_templates_do_not_retain_historical_results_or_images(self):
        added = [
            "parasite_6",
            "parasite_7",
            "elisa_mouse_4",
            "elisa_mouse_4_single",
            "elisa_mouse_7",
            "elisa_mouse_8_single",
            "elisa_rat_1",
            "pcr_9",
        ]
        for name in added:
            with self.subTest(template=name):
                path = TEMPLATES / f"{name}.docx"
                document = Document(path)
                self.assertEqual(document.tables[0].cell(1, 0).text, "{{dates}}")
                for table in document.tables[1:]:
                    reagent = "试剂盒" in table.cell(0, 0).text
                    for row in table.rows[1 if reagent else 0 :]:
                        for cell in row.cells[1:]:
                            self.assertFalse(cell.text.strip(), name)
                with ZipFile(path) as archive:
                    self.assertFalse(any(p.startswith(("word/media/", "word/comments")) for p in archive.namelist()))

    def test_all_templates_normalize_only_known_compatibility_issues_idempotently(self):
        for path in sorted(TEMPLATES.glob("*.docx")):
            with self.subTest(template=path.name), ZipFile(path) as archive:
                styles = archive.read("word/styles.xml")
                normalized_styles = normalize_style_order(styles)
                self.assertEqual(normalize_style_order(normalized_styles), normalized_styles)
                self.assertEqual(
                    [Counter(children) for children in style_children(normalized_styles)],
                    [Counter(children) for children in style_children(styles)],
                )
                assert_valid_style_order(self, normalized_styles)
                for name in archive.namelist():
                    if not name.startswith("word/footer") or not name.endswith(".xml"):
                        continue
                    footer = archive.read(name)
                    normalized_footer = normalize_footer_shape_ids(footer)
                    self.assertEqual(normalize_footer_shape_ids(normalized_footer), normalized_footer)
                    ids = footer_shape_ids(normalized_footer)
                    self.assertEqual(len(ids), len(set(ids)), name)
                    if footer_shape_ids(footer) == ids:
                        self.assertEqual(normalized_footer, footer, name)

    def test_footer_id_repair_leaves_later_unique_ids_unchanged(self):
        footer = f'''<?xml version="1.0" encoding="UTF-8"?>
        <w:ftr xmlns:w="{WORDPROCESSINGML}" xmlns:v="{VML}" xmlns:o="{OFFICE}">
          <w:p><w:r><w:pict>
            <v:shape id="_x0000_s1026" o:spid="_x0000_s1026"/>
            <v:shape id="_x0000_s1026" o:spid="_x0000_s1026"/>
            <v:shape id="_x0000_s1027" o:spid="_x0000_s1027"/>
          </w:pict></w:r></w:p>
        </w:ftr>'''.encode()
        normalized = normalize_footer_shape_ids(footer)
        self.assertEqual(footer_shape_ids(normalized), ["_x0000_s1026", "_x0000_s1028", "_x0000_s1027"])
        self.assertEqual(footer_shape_spids(normalized), ["_x0000_s1026", "_x0000_s1028", "_x0000_s1027"])
