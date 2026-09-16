import tempfile
import unittest
from base64 import b64decode
from pathlib import Path

from docx import Document
from docx.shared import Emu, Pt
from docx.table import Table

from server_app.domains.quarantine.document_layout import horizontal_result, image_table, sample_table

PNG = b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVQIHWP4z8DwHwAFgAI/ScL8KQAAAABJRU5ErkJggg==")


def row_height(row):
    return row.height.twips


def picture_extent(cell):
    extent = next(run._r.xpath(".//wp:extent")[0] for run in cell.paragraphs[0].runs if run._r.xpath(".//wp:extent"))
    return int(extent.get("cx")), int(extent.get("cy"))


class QuarantineDocumentLayoutTests(unittest.TestCase):
    def test_single_column_image_table_keeps_every_caption(self):
        document = Document()
        template = document.add_table(rows=4, cols=1)
        entries = [{"caption": f"图{index}"} for index in range(3)]
        result = Table(image_table(template._tbl, document, entries, Path(), []), document)
        self.assertEqual(len(result.columns), 1)
        self.assertEqual([row.cells[0].text for row in result.rows[1::2]], ["图0", "图1", "图2"])

    def test_result_chunks_keep_all_samples_and_separate_only_adjacent_tables(self):
        document = Document()
        template = document.add_table(rows=2, cols=5)
        samples = [{"id": str(i), "number": str(i)} for i in range(5)]
        project = {"sampleIds": [s["id"] for s in samples], "results": {}, "nc": "negative", "pc": "positive"}
        elements = horizontal_result(template._tbl, document, project, samples, controls=True)
        self.assertEqual([element.tag.rsplit("}", 1)[1] for element in elements], ["tbl", "p", "tbl", "p", "tbl"])
        actual = [cell.text for element in elements[::2] for cell in Table(element, document).rows[0].cells[1:-2]]
        self.assertEqual(actual, [str(i) for i in range(5)])

    def test_image_table_preserves_slot_widths_sizes_and_later_row_formats(self):
        document = Document()
        template = document.add_table(rows=4, cols=2)
        for row, height in zip(template.rows, (Pt(20), Pt(21), Pt(30), Pt(31)), strict=True):
            row.height = height
        for column, width in zip(template.columns, (Emu(2500000), Emu(1500000)), strict=True):
            for cell in column.cells:
                cell.width = width

        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            for index in range(4):
                (root / f"{index}.png").write_bytes(PNG)
            result = image_table(
                template._tbl,
                document,
                [{"caption": str(index), "attachment": {"storageName": f"{index}.png"}} for index in range(4)],
                root,
                [(1200000, 600000), (1200000, 700000), (500000, 800000), (900000, 400000)],
            )

        table = Table(result, document)
        for actual, expected in zip(
            [cell.width for cell in table.rows[0].cells], [Emu(2500000), Emu(1500000)], strict=True
        ):
            self.assertLess(abs(actual - expected), 200)
        self.assertEqual([row_height(row) for row in table.rows], [400, 420, 600, 620])
        self.assertEqual(picture_extent(table.cell(0, 0)), (600000, 600000))
        self.assertEqual(picture_extent(table.cell(2, 0)), (500000, 500000))

    def test_sample_table_uses_later_template_row_groups_before_reusing_the_last(self):
        document = Document()
        template = document.add_table(rows=4, cols=2)
        for row, height in zip(template.rows, (Pt(20), Pt(21), Pt(30), Pt(31)), strict=True):
            row.height = height

        samples = [
            {
                "id": str(index),
                "number": str(index),
                "sourceIds": [],
            }
            for index in range(3)
        ]
        result = sample_table(template._tbl, document, samples, {})
        table = Table(result, document)
        self.assertEqual([row_height(row) for row in table.rows], [400, 420, 600, 620, 600, 620])
