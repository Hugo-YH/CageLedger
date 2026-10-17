import io
import unittest

import openpyxl

from server_app.domains.billing.monthly_summary import (
    build_monthly_summary_rows,
    build_monthly_summary_xlsx,
    build_pi_reimbursement_context,
)


class MonthlyBillingSummaryTests(unittest.TestCase):
    def setUp(self):
        self.rooms = [
            {
                "id": "room-zj",
                "name": "8101",
                "facility": "zhujiang",
                "defaultSpecies": "mouse",
                "defaultBillingItem": "mouse_standard",
                "defaultCustomerType": "internal",
                "billingProfileConfigured": True,
                "billingProfileConfirmed": True,
            },
            {
                "id": "room-bwd",
                "name": "B201",
                "facility": "bioisland",
                "defaultSpecies": "mouse",
                "defaultBillingItem": "mouse_standard",
                "defaultCustomerType": "internal",
                "billingProfileConfigured": True,
                "billingProfileConfirmed": True,
            },
        ]
        self.applications = {
            "Z2026001": {
                "funding": "伦理经费-001",
                "projectStartDate": "2026-01-01",
                "projectEndDate": "2026-12-31",
            },
            "Z2026002": {
                "funding": "伦理经费-002",
                "projectStartDate": "2026-02-01",
                "projectEndDate": "2026-11-30",
            },
        }

    def test_splits_facilities_and_keeps_zero_amount_sheet(self):
        sheets = [
            sheet("sheet-zj", "Z2026001", "room-zj", "8101", 1),
            sheet("sheet-bwd", "Z2026001", "room-bwd", "B201", 1),
            sheet("sheet-zero", "Z2026002", "room-zj", "8101", 0),
        ]
        rows = build_monthly_summary_rows(
            "2026-06",
            sheets,
            self.rooms,
            self.applications,
            {"张教授": "independent"},
            {"张教授": {"fundBookNo": "本号-01", "reimbursementFormNo": "报销-01"}},
        )
        self.assertEqual(
            {(row["iacuc"], row["facility"]) for row in rows},
            {("Z2026001", "生物岛"), ("Z2026001", "珠江新城"), ("Z2026002", "珠江新城")},
        )
        self.assertTrue(all(row["amount"] == row["supportAmount"] + row["payableAmount"] for row in rows))
        rows_by_iacuc = {row["iacuc"]: row for row in rows if row["iacuc"] == "Z2026002"}
        self.assertEqual(next(row for row in rows if row["iacuc"] == "Z2026001")["funding"], "伦理经费-001")
        self.assertTrue(all(row["species"] == "小鼠" for row in rows))
        self.assertTrue(all(row["fundBookNo"] == "本号-01" for row in rows))
        self.assertEqual(rows_by_iacuc["Z2026002"]["amount"], 0)
        self.assertEqual(rows_by_iacuc["Z2026002"]["payableAmount"], 0)

    def test_workbook_matches_monthly_summary_columns_and_formula(self):
        rows = build_monthly_summary_rows(
            "2026-06",
            [sheet("sheet-zj", "Z2026001", "room-zj", "8101", 1)],
            self.rooms,
            self.applications,
            {"张教授": "independent"},
            {},
        )
        workbook = openpyxl.load_workbook(io.BytesIO(build_monthly_summary_xlsx("2026-06", rows)), data_only=False)
        worksheet = workbook.active
        self.assertEqual(worksheet["A1"].value, "2026年6月实验动物费用汇总")
        self.assertEqual(worksheet["D2"].value, "IACUC编号")
        self.assertEqual(worksheet["I3"].value, "=G3-H3")
        self.assertEqual(worksheet["G3"].number_format, "0.00")
        self.assertEqual(worksheet.freeze_panes, "A3")

    def test_invalid_sheet_error_identifies_pi_and_iacuc(self):
        invalid_sheet = sheet("invalid", "", "room-zj", "8101", 1)
        invalid_sheet["pi"] = ""
        with self.assertRaisesRegex(ValueError, r"项目负责人 未填写 / IACUC 未填写：缺少项目负责人"):
            build_monthly_summary_rows(
                "2026-06",
                [invalid_sheet],
                self.rooms,
                self.applications,
                {},
                {},
            )

    def test_expired_iacuc_before_month_marks_month_as_ineligible(self):
        applications = {
            **self.applications,
            "Z2026003": {"funding": "伦理经费-003", "projectEndDate": "2025-12-30"},
        }
        rows = build_monthly_summary_rows(
            "2026-06",
            [sheet("expired", "Z2026003", "room-zj", "8101", 1)],
            self.rooms,
            applications,
            {"张教授": "independent"},
            {},
        )
        self.assertEqual(rows[0]["notes"], "Z2026003 已于 2025-12-30 到期")

    def test_expired_iacuc_during_month_uses_next_day(self):
        applications = {
            **self.applications,
            "Z2026004": {"funding": "伦理经费-004", "projectEndDate": "2026-06-10"},
        }
        rows = build_monthly_summary_rows(
            "2026-06",
            [sheet("expired-in-month", "Z2026004", "room-zj", "8101", 1)],
            self.rooms,
            applications,
            {"张教授": "independent"},
            {},
        )
        self.assertEqual(rows[0]["notes"], "Z2026004 将于 2026-06-10 到期")

    def test_pi_reimbursement_context_uses_workflow_forms_only(self):
        workflows = [
            {
                "signedStatementReturned": True,
                "reimbursementForms": [
                    {"formNo": "报销-A", "amount": 120, "fundingBookNo": "本号-1"},
                    {"formNo": "报销-B", "amount": 80, "fundingBookNo": "本号-1"},
                ],
                "reimbursementFormNote": "已交回",
            }
        ]
        context = build_pi_reimbursement_context(workflows)
        self.assertEqual(context["statementReturned"], "已交回")
        self.assertEqual(context["fundBookNo"], "本号-1")
        self.assertEqual(context["reimbursementFormNo"], "报销-A、报销-B")
        self.assertEqual(context["reimbursementAmount"], 200)
        self.assertEqual(context["notes"], "已交回")

    def test_pi_reimbursement_context_empty_without_registered_forms(self):
        context = build_pi_reimbursement_context([])
        self.assertEqual(context["statementReturned"], "")
        self.assertEqual(context["fundBookNo"], "")
        self.assertEqual(context["reimbursementFormNo"], "")
        self.assertEqual(context["reimbursementAmount"], 0.0)
        self.assertEqual(context["notes"], "")

    def test_pi_reimbursement_context_keeps_statement_returned_without_forms(self):
        context = build_pi_reimbursement_context([{"signedStatementReturned": True}])
        self.assertEqual(context["statementReturned"], "已交回")
        self.assertEqual(context["fundBookNo"], "")
        self.assertEqual(context["reimbursementFormNo"], "")
        self.assertEqual(context["reimbursementAmount"], 0.0)
        self.assertEqual(context["notes"], "")

    def test_pi_reimbursement_context_merges_notes_across_workflows(self):
        context = build_pi_reimbursement_context(
            [
                {"reimbursementForms": [{"formNo": "报销-A", "amount": 10}], "reimbursementFormNote": "备注一"},
                {"reimbursementForms": [{"formNo": "报销-B", "amount": 20}], "reimbursementFormNote": "备注二"},
            ]
        )
        self.assertEqual(context["reimbursementFormNo"], "报销-A、报销-B")
        self.assertEqual(context["reimbursementAmount"], 30)
        self.assertEqual(context["notes"], "备注一；备注二")

    def test_workbook_includes_reimbursement_amount_and_notes_columns(self):
        rows = build_monthly_summary_rows(
            "2026-06",
            [sheet("sheet-zj", "Z2026001", "room-zj", "8101", 1)],
            self.rooms,
            self.applications,
            {"张教授": "independent"},
            {
                "张教授": {
                    "statementReturned": "已交回",
                    "fundBookNo": "本号-01",
                    "reimbursementFormNo": "报销-01",
                    "reimbursementAmount": 12.34,
                    "notes": "报销备注",
                }
            },
        )
        workbook = openpyxl.load_workbook(io.BytesIO(build_monthly_summary_xlsx("2026-06", rows)), data_only=False)
        worksheet = workbook.active
        self.assertEqual(worksheet["J2"].value, "结算单")
        self.assertEqual(worksheet["J3"].value, "已交回")
        self.assertEqual(worksheet["M2"].value, "报销金额（元）")
        self.assertEqual(worksheet["N2"].value, "报销备注")
        self.assertEqual(worksheet["M3"].value, 12.34)
        self.assertEqual(worksheet["N3"].value, "报销备注")
        self.assertEqual(worksheet["M3"].number_format, "0.00")
        self.assertEqual({str(item) for item in worksheet.merged_cells.ranges}, {"A1:S1"})
        self.assertEqual(worksheet.auto_filter.ref, "A2:S3")

    def test_workbook_preserves_complete_cell_style_matrix(self):
        rows = [
            {
                "pi": "张教授",
                "facility": "珠江新城",
                "iacuc": "Z2026001",
                "funding": "伦理经费-001",
                "species": "小鼠",
                "amount": 100.12,
                "supportAmount": 10.12,
                "payableAmount": 90,
                "statementReturned": "已交回",
                "fundBookNo": "本号-01",
                "reimbursementFormNo": "报销-01",
                "reimbursementAmount": 12.34,
                "reimbursementNotes": "报销备注",
                "projectStartDate": "2026-01-01",
                "projectEndDate": "2026-12-31",
                "notes": "备注",
                "roomNames": "8101",
                "technicalServiceAmount": 2.5,
            },
            {
                "pi": "李教授",
                "facility": "生物岛",
                "iacuc": "Z2026002",
                "funding": "伦理经费-002",
                "species": "大鼠",
                "amount": 200.24,
                "supportAmount": 20.24,
                "payableAmount": 180,
                "statementReturned": "",
                "fundBookNo": "",
                "reimbursementFormNo": "",
                "reimbursementAmount": 0,
                "reimbursementNotes": "",
                "projectStartDate": "2026-02-01",
                "projectEndDate": "2026-11-30",
                "notes": "",
                "roomNames": "B201",
                "technicalServiceAmount": 0,
            },
        ]
        workbook = openpyxl.load_workbook(io.BytesIO(build_monthly_summary_xlsx("2026-06", rows)), data_only=False)
        worksheet = workbook.active
        wrapped_columns = {5, 6, 14, 17, 18}
        numeric_columns = {7, 8, 9, 13, 19}

        self.assertEqual(worksheet["A1"].font.name, "Microsoft YaHei")
        self.assertEqual(worksheet["A1"].font.sz, 14)
        self.assertTrue(worksheet["A1"].font.bold)
        self.assertEqual(worksheet["A1"].alignment.horizontal, "center")
        self.assertEqual(worksheet["A1"].alignment.vertical, "center")
        for column in range(1, 20):
            header = worksheet.cell(2, column)
            self.assertEqual(header.font.name, "Microsoft YaHei")
            self.assertEqual(header.font.sz, 10)
            self.assertTrue(header.font.bold)
            self.assertEqual(header.fill.fill_type, "solid")
            self.assertEqual(header.fill.fgColor.rgb, "00D9EAD3")
            self.assertTrue(header.alignment.wrap_text)
            self.assertEqual(header.alignment.horizontal, "center")
            self.assertEqual(header.alignment.vertical, "center")
            self.assertEqual(header.border.left.style, "thin")
            self.assertEqual(header.border.right.style, "thin")
            self.assertEqual(header.border.top.style, "thin")
            self.assertEqual(header.border.bottom.style, "thin")
            for row in range(3, 5):
                cell = worksheet.cell(row, column)
                self.assertEqual(cell.font.name, "Microsoft YaHei")
                self.assertEqual(cell.font.sz, 10)
                self.assertFalse(cell.font.bold)
                self.assertEqual(cell.alignment.horizontal, "center")
                self.assertEqual(cell.alignment.vertical, "center")
                self.assertEqual(bool(cell.alignment.wrap_text), column in wrapped_columns)
                self.assertEqual(cell.number_format, "0.00" if column in numeric_columns else "General")
                self.assertEqual(cell.border.left.style, "thin")
                self.assertEqual(cell.border.right.style, "thin")
                self.assertEqual(cell.border.top.style, "thin")
                self.assertEqual(cell.border.bottom.style, "thin")
        self.assertEqual(worksheet["I3"].data_type, "f")
        self.assertEqual(worksheet["I4"].data_type, "f")
        self.assertEqual(worksheet["A3"].data_type, "n")
        self.assertEqual(worksheet["A4"].data_type, "n")
        self.assertEqual(worksheet.row_dimensions[1].height, 28)
        self.assertEqual(worksheet.row_dimensions[2].height, 36)
        self.assertEqual(worksheet.row_dimensions[3].height, 24)
        self.assertEqual(worksheet.row_dimensions[4].height, 24)
        self.assertEqual(worksheet.column_dimensions["A"].width, 9)
        self.assertEqual(worksheet.column_dimensions["S"].width, 19)
        self.assertEqual({str(item) for item in worksheet.merged_cells.ranges}, {"A1:S1"})
        self.assertEqual(worksheet.freeze_panes, "A3")
        self.assertEqual(worksheet.auto_filter.ref, "A2:S4")
        self.assertTrue(workbook.calculation.fullCalcOnLoad)
        self.assertTrue(workbook.calculation.forceFullCalc)

    def test_empty_workbook_keeps_the_original_two_row_shape(self):
        workbook = openpyxl.load_workbook(io.BytesIO(build_monthly_summary_xlsx("2026-06", [])), data_only=False)
        worksheet = workbook.active

        self.assertEqual(worksheet.max_row, 2)
        self.assertEqual(worksheet.max_column, 19)
        self.assertEqual(worksheet.calculate_dimension(), "A1:S2")
        self.assertEqual(worksheet.auto_filter.ref, "A2:S2")
        self.assertEqual(worksheet.row_dimensions[1].height, 28)
        self.assertEqual(worksheet.row_dimensions[2].height, 36)
        self.assertNotIn(3, worksheet.row_dimensions)
        self.assertEqual({str(item) for item in worksheet.merged_cells.ranges}, {"A1:S1"})
        self.assertEqual(worksheet.freeze_panes, "A3")


def sheet(sheet_id, iacuc, room_id, room_name, cages):
    return {
        "id": sheet_id,
        "month": "2026-06",
        "iacuc": iacuc,
        "pi": "张教授",
        "roomId": room_id,
        "roomName": room_name,
        "initialCageCount": cages,
        "initialAnimalCount": 0,
        "rows": [],
    }


if __name__ == "__main__":
    unittest.main()
