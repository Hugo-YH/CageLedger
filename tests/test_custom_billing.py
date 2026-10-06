import io
import json
import sqlite3
import unittest
from contextlib import closing

from openpyxl import load_workbook

import server
from server_app.cache import invalidate_data_cache_prefixes
from server_app.domains.billing.custom_billing import (
    normalize_custom_billing_segments,
    validate_custom_billing_segments,
)
from server_app.domains.billing.settlement_xlsx import build_settlement_workbook
from server_app.pdf.documents import billing_statement_html
from server_app.shared.concurrency import StaleWriteError


class AllQuantityCustomBillingTests(unittest.TestCase):
    def setUp(self):
        invalidate_data_cache_prefixes("quantity_sheets::")
        self.addCleanup(invalidate_data_cache_prefixes, "quantity_sheets::")

    def fixture(self, animal=True):
        unit = "animal_day" if animal else "cage_day"
        count_key = "animalCount" if animal else "cageCount"
        sheet = {
            "id": "custom-all",
            "month": "2026-07",
            "roomId": "custom-room",
            "iacuc": "Z-ALL",
            "pi": "测试负责人",
            "billingUnit": unit,
            "initialAnimalCount": 10,
            "initialCageCount": 10,
            "fullExemption": True,
            "rows": [
                {"id": f"row-{day}", "date": f"2026-07-{day:02d}", count_key: count}
                for day, count in [(1, 10), (2, 12), (3, 7), (4, 0)]
            ],
            "customBillingSegments": [
                {
                    "id": "all",
                    "startDate": "2026-07-01",
                    "endDate": "2026-07-03",
                    "quantityMode": "all",
                    "quantity": None,
                    "unitPrice": 3,
                    "note": "特殊饲养",
                }
            ],
        }
        room = {
            "id": "custom-room",
            "name": "测试房间",
            "defaultBillingItem": "rabbit" if animal else "mouse_standard",
            "defaultCustomerType": "internal",
            "defaultSpecies": "rabbit" if animal else "mouse",
            "billingProfileConfigured": True,
            "billingProfileConfirmed": True,
        }
        return sheet, room

    def test_all_follows_each_days_animals_and_cages_without_free_allowance_or_tiers(self):
        for animal in (True, False):
            with self.subTest(animal=animal):
                sheet, room = self.fixture(animal)
                lines = server.quantity_sheet_statement_lines([sheet], 100, [room], {})
                self.assertEqual([line["amount"] for line in lines[:4]], [30, 36, 21, 0])
                self.assertEqual(sum(line["amount"] for line in lines), 87)
                key = "animalCount" if animal else "cageCount"
                for line, count in zip(lines[:3], [10, 12, 7], strict=True):
                    custom = next(item for item in line["iacucBreakdown"] if item["customBilling"])
                    self.assertEqual(custom[key], count)
                    self.assertEqual(custom["customBillingQuantityMode"], "all")
                    self.assertEqual(custom["freeCages"], 0)
                    self.assertFalse(custom["tiered"])
                    self.assertFalse(custom["fullExemption"])

    def test_dates_outside_interval_use_standard_price(self):
        sheet, room = self.fixture()
        sheet["fullExemption"] = False
        sheet["customBillingSegments"][0].update(startDate="2026-07-02", endDate="2026-07-02")
        lines = server.quantity_sheet_statement_lines([sheet], 0, [room], {})
        self.assertEqual([line["amount"] for line in lines[:4]], [50, 36, 35, 0])

    def test_full_quantity_overlaps_are_rejected_even_on_empty_days(self):
        for mode in ("fixed", "all"):
            sheet, room = self.fixture()
            sheet["customBillingSegments"][0].update(endDate="2026-07-04")
            sheet["customBillingSegments"].append(
                {
                    "id": "overlap",
                    "startDate": "2026-07-04",
                    "endDate": "2026-07-04",
                    "quantityMode": mode,
                    "quantity": 1 if mode == "fixed" else None,
                    "unitPrice": 4,
                }
            )
            with self.subTest(mode=mode), self.assertRaisesRegex(ValueError, "区间不能.*重叠"):
                validate_custom_billing_segments([sheet], [room])

    def test_blank_fixed_quantity_is_not_all_and_legacy_null_is_all(self):
        sheet, room = self.fixture()
        sheet["customBillingSegments"][0]["quantityMode"] = "fixed"
        with self.assertRaisesRegex(ValueError, "每日适用数量"):
            validate_custom_billing_segments([sheet], [room])
        del sheet["customBillingSegments"][0]["quantityMode"]
        normalized = normalize_custom_billing_segments(sheet, sheet["month"])
        self.assertEqual(normalized[0]["quantityMode"], "all")
        validate_custom_billing_segments([sheet], [room])

    def test_sqlite_roundtrip_audit_and_stale_save_keep_all_quantity_mode(self):
        sheet, room = self.fixture()
        actor = {"id": "u-custom", "username": "admin", "displayName": "测试管理员", "role": "admin", "roomIds": []}
        with closing(sqlite3.connect(":memory:")) as conn:
            conn.row_factory = sqlite3.Row
            server.initialize_schema(conn)
            conn.execute(
                "INSERT INTO rooms (id, name, payload) VALUES (?, ?, ?)", (room["id"], room["name"], json.dumps(room))
            )
            saved, _, _, _, status, _ = server.save_quantity_sheet(conn, {"sheet": sheet}, actor)
            self.assertEqual(status, 201)
            payload = json.loads(
                conn.execute("SELECT payload FROM quantity_sheets WHERE id = ?", (sheet["id"],)).fetchone()[0]
            )
            self.assertEqual(payload["customBillingSegments"][0]["quantityMode"], "all")
            self.assertIsNone(payload["customBillingSegments"][0]["quantity"])
            self.assertEqual(
                conn.execute(
                    "SELECT count(*) FROM audit_events WHERE entity_id = ? AND action = 'quantity_sheet.created'",
                    (sheet["id"],),
                ).fetchone()[0],
                1,
            )
            conn.commit()
            invalidate_data_cache_prefixes("quantity_sheets::")
            with self.assertRaisesRegex(StaleWriteError, "数量统计表"):
                server.save_quantity_sheet(conn, {"sheet": saved, "expectedUpdatedAt": "stale"}, actor, sheet["id"])

    def test_export_uses_daily_quantities_and_does_not_claim_a_fixed_daily_count(self):
        sheet, room = self.fixture()
        lines = server.quantity_sheet_statement_lines([sheet], 0, [room], {})
        statement = {
            "id": "custom-statement",
            "month": sheet["month"],
            "pi": sheet["pi"],
            "iacucs": [sheet["iacuc"]],
            "sourceType": "quantity_sheet",
            "totalAmount": 87,
        }
        html = billing_statement_html(statement, lines)
        self.assertIn("按每日实际结余只数计费", html)
        self.assertNotIn("每日10只", html)
        workbook = load_workbook(io.BytesIO(build_settlement_workbook([(statement, lines)])))
        ws = workbook.active
        rows = [
            row
            for row in ws.iter_rows()
            if len(row) > 5 and row[3].value in (10, 12, 7) and row[4].value == 3 and row[5].data_type == "f"
        ]
        self.assertEqual([row[3].value for row in rows], [10, 12, 7])
