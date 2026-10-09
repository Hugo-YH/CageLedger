import copy
import unittest
from io import BytesIO
from unittest.mock import patch

from docx import Document
from test_quarantine_document_templates import report_snapshot

from server_app.domains.quarantine import summary


class QuarantineSummaryTests(unittest.TestCase):
    def test_custom_projects_are_not_assigned_a_biological_category(self):
        self.assertEqual(summary.category("pcr", "自定义病毒项目"), 3)
        self.assertEqual(summary.category("elisa_mouse", "自定义项目"), 3)

    def test_counts_are_assay_groups_and_preserve_incomplete_and_suspect_results(self):
        test = report_snapshot("elisa_mouse", 5)["test"]
        project = test["projects"][0]
        project["results"] = dict(
            zip(project["sampleIds"], ["positive", "negative", "suspect", "not_tested", ""], strict=True)
        )
        for sample in test["samples"]:
            sample["poolCount"] = 10
        result, notes = summary.project_result(test, project)
        self.assertEqual(result, "1/3")
        self.assertEqual(notes, "1组可疑；1组未检测；1组未填写")
        project["results"] = {sid: "not_tested" for sid in project["sampleIds"]}
        self.assertEqual(summary.project_result(test, project), ("—", "5组未检测"))

    def test_distinct_assays_do_not_fill_ambiguous_historical_rows(self):
        records = [
            {"test": report_snapshot(method, 2)["test"], "issued": False, "number": ""}
            for method in ("parasite", "elisa_mouse", "pcr")
        ]
        document = Document(
            BytesIO(summary.generate({"records": records, "batch": {"sources": [{"species": "小鼠"}]}}))
        )
        rows = {row.cells[0].text.strip(): row for row in document.tables[0].rows}
        for name in ["蠕虫（All helminths）", "小鼠细小病毒 (Mouse Parvovirus, MPV)"]:
            self.assertEqual(rows[name].cells[4].text, "—")
        for name in ["体内寄生虫", "小鼠微小病毒抗体（MVM）", "嗜肺巴斯德杆菌H型", "嗜肺巴斯德杆菌J型"]:
            self.assertEqual(rows[name].cells[4].text, "0/2")
        self.assertIn("草稿", document.sections[0].footer.paragraphs[0].text)
        self.assertIn("混样不折算", "\n".join(p.text for p in document.paragraphs))

    def test_issued_snapshot_is_used_and_retests_remain_separate_from_correction(self):
        original = {"id": "original", "testDate": "2026-09-02", "method": "parasite", "state": "issued"}
        correction = {**original, "id": "corrected", "correctionOf": "original"}
        retest = {**original, "id": "retest", "retestOf": "corrected"}
        frozen = {**correction, "conclusion": "冻结内容"}
        report = {"version": 1, "number": "Q-1", "snapshot": {"test": frozen}}

        def items(conn, entity, parent):
            if entity == "tests":
                return [original, correction, retest]
            return [report] if parent == "corrected" else []

        with (
            patch.object(summary.repo, "get", return_value={"sources": []}),
            patch.object(summary.repo, "all_items", side_effect=items),
        ):
            result = summary.batch_snapshot(None, "batch")
        self.assertEqual([r["test"]["id"] for r in result["records"]], ["corrected", "retest"])
        self.assertEqual(result["records"][0]["test"]["conclusion"], "冻结内容")
        self.assertTrue(result["records"][0]["issued"])
        self.assertFalse(result["records"][1]["issued"])

    def test_self_title_and_mixed_dates_and_retest_notes_are_explicit(self):
        test = report_snapshot("parasite", 1)["test"]
        repeat = copy.deepcopy(test)
        repeat.update(testDate="2026-09-10", retestOf="previous")
        data = {
            "batch": {"sources": []},
            "records": [{"test": t, "issued": True, "number": "Q-1"} for t in (test, repeat)],
        }
        document = Document(BytesIO(summary.generate(data, kind="self")))
        self.assertIn("实验动物自检检测报告", [p.text for p in document.paragraphs])
        row = next(r for r in document.tables[0].rows if r.cells[0].text == "体内寄生虫")
        self.assertEqual(row.cells[4].text, "2026-09-03\n0/1\n2026-09-10\n复检 0/1")
        self.assertIn("复检", row.cells[5].text)
        self.assertNotIn("草稿", document.sections[0].footer.paragraphs[0].text)
        self.assertIn("日期：", next(p.text for p in document.paragraphs if p.text.startswith("检测人：")))
        with self.assertRaisesRegex(ValueError, "类型无效"):
            summary.generate(data, kind="bad")
