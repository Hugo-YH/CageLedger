import copy
import sqlite3
import unittest
from io import BytesIO
from unittest.mock import Mock, patch
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

from server_app.domains.intake import summary
from server_app.repositories.entities import upsert_intake_batch
from server_app.web import intake_summary


def reservation(index=1, **overrides):
    return {
        "id": f"summary-{index}",
        "batchNo": f"（Z2026001）202610{index:04d}",
        "purchaseOrderNo": f"PO-{index}",
        "iacuc": "Z2026001",
        "owner": "预约人",
        "pi": "项目负责人",
        "supplier": "实验供应单位",
        "strainRaw": "C57 <原文> & ICR",
        "strainStandard": "标准化品系",
        "species": "mouse",
        "sex": "雌雄各半",
        "quantity": 6,
        "finalCardCount": 2,
        "intakeDate": "2026-10-12",
        "roomName": "8105",
        "status": "pending_print",
        "updatedAt": "2026-10-10T10:00:00",
        "husbandryDays": 30,
        "notes": "请核对耳标\n保持分笼",
        "rawMessage": "预约原文含8周龄，20–22克\n下午送达",
        **overrides,
    }


def paragraphs(document):
    return [Paragraph(element, document) for element in document.element.body.xpath(".//w:p")]


def document_text(document):
    return "\n".join(p.text for p in paragraphs(document))


class IntakeSummaryTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        self.conn.row_factory = sqlite3.Row
        self.conn.execute("""CREATE TABLE intake_batches (
            id TEXT PRIMARY KEY, batch_no TEXT, iacuc TEXT, supplier TEXT, pi TEXT, owner TEXT,
            quantity INTEGER, card_count INTEGER, room_name TEXT, intake_date TEXT,
            status TEXT, updated_at TEXT, payload TEXT)""")
        self.conn.execute("CREATE INDEX idx_intake_date ON intake_batches(intake_date)")
        self.addCleanup(self.conn.close)

    def test_interval_includes_both_days_and_all_statuses_without_page_limit_or_writes(self):
        for index in range(125):
            upsert_intake_batch(
                self.conn, reservation(index, status=["draft", "pending_print", "printed", "received"][index % 4])
            )
        upsert_intake_batch(self.conn, reservation(200, intakeDate="2026-10-13", roomName="8103"))
        for index, day in [(201, "2026-10-11"), (202, "2026-10-14"), (203, "")]:
            upsert_intake_batch(self.conn, reservation(index, intakeDate=day))
        changes = self.conn.total_changes
        before = [tuple(row) for row in self.conn.execute("SELECT * FROM intake_batches")]
        snapshot = summary.snapshot(self.conn, "2026-10-12", "2026-10-13")
        self.assertEqual(len(snapshot["batches"]), 126)
        self.assertEqual(snapshot["batches"][-1]["id"], "summary-200")
        summary.generate(snapshot)
        self.assertEqual(changes, self.conn.total_changes)
        self.assertEqual(before, [tuple(row) for row in self.conn.execute("SELECT * FROM intake_batches")])
        plan = self.conn.execute(
            "EXPLAIN QUERY PLAN SELECT payload FROM intake_batches WHERE intake_date >= ? AND intake_date <= ?",
            ("2026-10-12", "2026-10-13"),
        ).fetchall()
        self.assertIn("idx_intake_date", " ".join(row["detail"] for row in plan))

    def test_invalid_or_empty_intervals_fail_before_generating_a_blank_document(self):
        for start, end in [
            ("", "2026-10-12"),
            ("20261012", "2026-10-13"),
            ("2026-02-30", "2026-10-13"),
            ("2026-10-13", "2026-10-12"),
            ("2026-10-12' OR 1=1", "2026-10-13"),
        ]:
            with self.subTest(start=start), self.assertRaises(ValueError):
                summary.snapshot(self.conn, start, end)
        with self.assertRaisesRegex(ValueError, "没有预约接收批次"):
            summary.snapshot(self.conn, "2026-10-12", "2026-10-13")

    def test_word_preserves_reservation_details_and_blanks_and_groups_by_date_room(self):
        entries = [
            reservation(3, intakeDate="2026-10-13"),
            reservation(2, roomName="8105", status="received"),
            reservation(1, roomName="8103", quantity=0, husbandryDays=0),
        ]
        original = copy.deepcopy(entries)
        document = Document(
            BytesIO(summary.generate({"startDate": "2026-10-12", "endDate": "2026-10-13", "batches": entries}))
        )
        content = document_text(document)
        for value in [
            "共 3 批",
            "预约人（项目负责人）",
            "C57 <原文> & ICR",
            "数量：0",
            "Z2026001",
            "｜已接收｜",
        ]:
            self.assertIn(value, content)
        self.assertNotIn("标准化品系", content)
        self.assertIn("｜未打印｜", content)
        for removed in (
            "预约原文：",
            "8周龄",
            "签收：",
            "实收数量：",
            "IACUC",
            "笼数",
            "动物种类",
            "性别",
            "饲养天数",
            "预约备注",
            "请核对耳标",
        ):
            self.assertNotIn(removed, content)
        self.assertNotIn("✔", content)
        self.assertLess(content.index("PO-1"), content.index("PO-2"))
        self.assertLess(content.index("PO-2"), content.index("PO-3"))
        self.assertEqual(entries, original)
        self.assertEqual(len(document.tables), 3)
        self.assertAlmostEqual(document.sections[0].page_width.mm, 210, places=1)
        self.assertAlmostEqual(document.sections[0].page_height.mm, 297, places=1)
        self.assertEqual(document.sections[0]._sectPr.find(qn("w:docGrid")).get(qn("w:type")), "default")
        self.assertEqual(document.styles["Normal"].element.rPr.rFonts.get(qn("w:eastAsia")), "微软雅黑")
        for p in document.paragraphs:
            if p.text.startswith("□ "):
                self.assertIn(p.paragraph_format.left_indent, (None, 0))
                self.assertTrue(p.paragraph_format.keep_with_next)
                self.assertTrue(p.paragraph_format.keep_together)
                self.assertEqual(p._p.pPr.find(qn("w:shd")).get(qn("w:fill")), "F3D8C4")
                self.assertFalse(p._p.xpath(".//w:rPr/w:shd"))
                tab = p._p.pPr.find(qn("w:tabs"))[0]
                self.assertEqual(tab.get(qn("w:val")), "right")
                self.assertAlmostEqual(int(tab.get(qn("w:pos"))), 10205, delta=2)
                self.assertTrue(p.runs[-1].text.startswith("\t｜"))
        self.assertEqual([p.text for p in document.paragraphs].count("8103"), 1)
        room = next(p for p in document.paragraphs if p.text == "8103")
        self.assertIn(room.paragraph_format.left_indent, (None, 0))
        self.assertEqual(room._p.pPr.find(qn("w:shd")).get(qn("w:fill")), "CED1E9")
        self.assertNotIn("饲养房间", content)

    def test_detail_rows_use_fixed_half_and_third_widths_without_borders_or_fixed_heights(self):
        item = reservation(supplier="很长的供应商名称" * 8, strainRaw="长品系名称" * 8)
        document = Document(
            BytesIO(summary.generate({"startDate": "2026-10-12", "endDate": "2026-10-12", "batches": [item]}))
        )
        table = document.tables[0]
        self.assertFalse(table.autofit)
        indent = int(table._tbl.tblPr.find(qn("w:tblInd")).get(qn("w:w")))
        width = int(table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w")))
        self.assertAlmostEqual(indent, 283, delta=1)
        self.assertAlmostEqual(width, 9921, delta=1)
        self.assertAlmostEqual(indent + width, 10205, delta=2)
        self.assertEqual([len(row._tr.tc_lst) for row in table.rows], [2, 3])
        for row, span in zip(table.rows, (3, 2), strict=True):
            widths = [int(cell.tcPr.tcW.get(qn("w:w"))) for cell in row._tr.tc_lst]
            self.assertEqual(len(set(widths)), 1)
            self.assertAlmostEqual(sum(widths), int(table._tbl.tblPr.find(qn("w:tblW")).get(qn("w:w"))), delta=2)
            self.assertTrue(all(cell.tcPr.gridSpan.val == span for cell in row._tr.tc_lst))
            self.assertFalse(row._tr.xpath("./w:trPr/w:trHeight"))
        self.assertTrue(all(edge.get(qn("w:val")) == "nil" for edge in table._tbl.tblPr.find(qn("w:tblBorders"))))
        self.assertIn(item["supplier"], document_text(document))
        self.assertIn(item["strainRaw"], document_text(document))

    def test_missing_fields_are_not_invented_and_xml_control_characters_are_removed(self):
        item = reservation(
            owner="",
            pi="",
            quantity=None,
            strainRaw="",
            strainStandard="备用品系",
            supplier="合法\x00文本\x0b\ud800",
            rawMessage="",
            finalCardCount=None,
        )
        document = Document(
            BytesIO(summary.generate({"startDate": "2026-10-12", "endDate": "2026-10-12", "batches": [item]}))
        )
        content = document_text(document)
        for value in ["负责人未填写", "数量：未填写", "备用品系", "合法文本"]:
            self.assertIn(value, content)
        self.assertNotIn("预约原文：", content)

    def test_field_labels_are_bold_but_values_with_colons_and_line_breaks_stay_normal(self):
        item = reservation(
            purchaseOrderNo="PO202610005",
            strainRaw="品系：A <B>",
            supplier="供应商：按预约原样记录\n第二行供应商",
            project="示例项目",
            receiverName="示例接收人",
            vetPhone="020-12345678",
            notes="供应商：按预约原样记录\n第二行备注",
        )
        document = Document(
            BytesIO(summary.generate({"startDate": "2026-10-12", "endDate": "2026-10-12", "batches": [item]}))
        )
        details = [p for p in paragraphs(document) if any(run.bold is False for run in p.runs)]
        labels = {run.text for p in details for run in p.runs if run.bold}
        self.assertEqual(
            labels,
            {
                "采购单号",
                "批次号",
                "供应商",
                "品系",
                "数量",
            },
        )
        for value in ("：PO202610005", "：品系：A <B>", "：供应商：按预约原样记录\n第二行供应商"):
            run = next(run for p in details for run in p.runs if run.text == value)
            self.assertFalse(run.bold)
        for removed in ("项目名称", "示例项目", "接收人", "示例接收人", "兽医电话", "020-12345678", "饲养房间"):
            self.assertNotIn(removed, document_text(document))

    def test_template_contains_only_blank_layout_and_all_status_tags_are_exported(self):
        template = Document(summary.TEMPLATE)
        self.assertFalse(template.paragraphs)
        self.assertEqual(template.core_properties.author, "CageLedger")
        with ZipFile(summary.TEMPLATE) as archive:
            self.assertFalse(
                any("media/" in name or "custom.xml" in name or "comments" in name for name in archive.namelist())
            )
        entries = [reservation(i, status=status) for i, status in enumerate(summary.STATUS_LABELS)]
        document = Document(
            BytesIO(summary.generate({"startDate": "2026-10-12", "endDate": "2026-10-12", "batches": entries}))
        )
        content = document_text(document)
        for label in summary.STATUS_LABELS.values():
            self.assertIn(f"｜{label}｜", content)

    def test_handler_rejects_anonymous_and_allows_list_roles_with_private_word_download(self):
        handler = Mock(path="/api/intake-batches/summary.docx?startDate=2026-10-12&endDate=2026-10-12")
        handler.require_user.return_value = None
        with patch.object(intake_summary, "connect_db") as connect:
            self.assertTrue(intake_summary.handle(handler, "/api/intake-batches/summary.docx"))
            connect.assert_not_called()
        handler.send_download.assert_not_called()
        upsert_intake_batch(self.conn, reservation())
        for role in ("admin", "room_admin"):
            handler.require_user.return_value = {"role": role, "roomIds": []}
            with patch.object(intake_summary, "connect_db") as connect:
                connect.return_value.__enter__.return_value = self.conn
                self.assertTrue(intake_summary.handle(handler, "/api/intake-batches/summary.docx"))
            data, name, mime = handler.send_download.call_args.args
            self.assertEqual(name, "接收汇总_2026-10-12_2026-10-12.docx")
            self.assertEqual(mime, summary.DOCX_MIME)
            self.assertTrue(Document(BytesIO(data)).paragraphs)
        handler.path = "/api/intake-batches/summary.docx?startDate=bad"
        self.assertTrue(intake_summary.handle(handler, "/api/intake-batches/summary.docx"))
        self.assertEqual(handler.send_json.call_args.args[1], 400)
        self.assertFalse(intake_summary.handle(handler, "/api/intake-batches"))


if __name__ == "__main__":
    unittest.main()
