import io
import sqlite3
import tempfile
import threading
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl

from server_app.domains.reimbursement import application
from server_app.domains.reimbursement.application import (
    import_arrears_reimbursement_workbook,
    import_monthly_reimbursement_workbook,
)
from server_app.domains.reimbursement.facade import get_reimbursement_record, upsert_reimbursement_record
from server_app.domains.reimbursement.importer import build_monthly_import_groups, month_key
from server_app.legacy import initialize_schema
from server_app.web import reimbursement_actions

ACTOR = {"id": "admin", "username": "admin", "displayName": "测试管理员", "role": "admin"}
HEADERS = ["项目负责人", "IACUC", "产生的饲养费", "单位支持", "应缴纳"]


def monthly_book(title="2026-10", amount=12.34):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = title
    sheet.append(HEADERS)
    sheet.append(["合成负责人", "Z2026001", amount, 0, amount])
    return book


def workbook_bytes(book):
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


class ReimbursementImportValidationTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        self.conn.row_factory = sqlite3.Row
        initialize_schema(self.conn)
        self.conn.commit()
        self.addCleanup(self.conn.close)

    def test_months_are_calendar_months_and_supported_legacy_formats_remain(self):
        for value in ("2026-00", "2026-13", "202613", "2026年13月", "2026-123", 202600, 202613, 202613.0):
            with self.subTest(value=value):
                self.assertEqual(month_key(value), "")
        for value in ("202610", "2026-10", "2026/10", "2026.10", "2026年10月", "2026-10-31", 202610, 202610.0):
            self.assertEqual(month_key(value), "2026-10")
        self.assertEqual(month_key(45000), "2023-03")

    def test_invalid_month_rejects_entire_workbook_before_any_write(self):
        book = monthly_book()
        other = book.create_sheet("2026-13")
        other.append(HEADERS)
        other.append(["第二负责人", "Z2026002", 10, 0, 10])
        before = self.conn.total_changes
        with self.assertRaisesRegex(ValueError, "月份"):
            import_monthly_reimbursement_workbook(self.conn, workbook_bytes(book), ACTOR)
        self.assertEqual(self.conn.total_changes, before)

    def test_invalid_money_reports_cell_and_never_partially_writes(self):
        for value in ("abc", "NaN", "Infinity", "1,2", True):
            for column in (3, 4, 5):
                with self.subTest(value=value, column=column):
                    book = monthly_book()
                    book.active.append(["第二负责人", "Z2026002", 10, 0, 10])
                    book.active.cell(3, column).value = value
                    before = self.conn.total_changes
                    coordinate = book.active.cell(3, column).coordinate
                    with self.assertRaisesRegex(ValueError, f"2026-10.*{coordinate}"):
                        import_monthly_reimbursement_workbook(self.conn, workbook_bytes(book), ACTOR)
                    self.assertEqual(self.conn.total_changes, before)

    def test_formatted_money_and_empty_optional_amounts_remain_compatible(self):
        book = monthly_book(amount="1,234.56")
        book.active.cell(2, 4).value = None
        result = import_monthly_reimbursement_workbook(self.conn, workbook_bytes(book), ACTOR)
        self.assertEqual(result["count"], 1)
        self.assertEqual(result["items"][0]["currentMonthAmount"], 1234.56)
        self.assertEqual(result["items"][0]["payableAmount"], 1234.56)
        repeated = import_monthly_reimbursement_workbook(self.conn, workbook_bytes(book), ACTOR)
        self.assertEqual(result["items"][0]["id"], repeated["items"][0]["id"])
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM reimbursement_records").fetchone()[0], 1)

    def test_empty_markers_and_existing_rounding_are_preserved(self):
        for marker in (None, "", "-", "—"):
            book = monthly_book(amount=2.675)
            book.active.cell(2, 4).value = marker
            result = build_monthly_import_groups(book)["2026-10|合成负责人"]
            self.assertEqual(result["currentMonthAmount"], 2.67)
            self.assertEqual(result["supportAmount"], 0)

    def test_empty_gap_does_not_hide_later_records_and_row_limit_is_explicit(self):
        book = monthly_book()
        book.active.append([])
        for row in (23, 1200):
            book.active.cell(row, 1).value = f"合成负责人{row}"
            book.active.cell(row, 3).value = 1
            book.active.cell(row, 5).value = 1
        self.assertEqual(len(build_monthly_import_groups(book)), 3)
        book.active.cell(1201, 1).value = "超过范围"
        with self.assertRaisesRegex(ValueError, "1200"):
            build_monthly_import_groups(book)

    def test_candidate_sheet_missing_amount_header_is_not_silently_skipped(self):
        book = monthly_book()
        other = book.create_sheet("2026-11")
        other.append(["项目负责人", "IACUC", "应缴纳"])
        other.append(["第二负责人", "Z2026002", 10])
        with self.assertRaisesRegex(ValueError, "2026-11.*表头"):
            build_monthly_import_groups(book)

    def test_arrears_validate_all_rows_and_months_before_writing(self):
        for invalid_month in (False, True):
            book = openpyxl.Workbook()
            book.active.append(
                [
                    "序号",
                    "项目负责人",
                    "动物品系",
                    "经费",
                    "累计欠缴",
                    "2026-10",
                    "2026-13" if invalid_month else "2026-11",
                ]
            )
            book.active.append([1, "第一负责人", "小鼠", "", 10, 10, 0])
            book.active.append([2, "第二负责人", "小鼠", "", 20, 10, "abc"])
            before = self.conn.total_changes
            with self.assertRaisesRegex(ValueError, "月份" if invalid_month else "G3"):
                import_arrears_reimbursement_workbook(self.conn, workbook_bytes(book), ACTOR)
            self.assertEqual(self.conn.total_changes, before)

    def test_numeric_month_headers_and_columns_beyond_limit_cannot_silently_misimport(self):
        for month, error in ((202610, None), (202613, "月份")):
            book = openpyxl.Workbook()
            book.active.append(["序号", "项目负责人", "动物品系", "经费", "累计欠缴", month])
            book.active.append([1, "合成负责人", "小鼠", "", 10, 10])
            if error:
                before = self.conn.total_changes
                with self.assertRaisesRegex(ValueError, error):
                    import_arrears_reimbursement_workbook(self.conn, workbook_bytes(book), ACTOR)
                self.assertEqual(self.conn.total_changes, before)
            else:
                result = import_arrears_reimbursement_workbook(self.conn, workbook_bytes(book), ACTOR)
                self.assertEqual(result["items"][0]["month"], "2026-10")
        book.active.cell(1, 121).value = "2026-11"
        book.active.cell(1, 6).value = "2026-10"
        book.active.cell(2, 121).value = 10
        before = self.conn.total_changes
        with self.assertRaisesRegex(ValueError, "120列"):
            import_arrears_reimbursement_workbook(self.conn, workbook_bytes(book), ACTOR)
        self.assertEqual(self.conn.total_changes, before)

    def test_import_http_error_is_400_and_admin_boundary_is_unchanged(self):
        class Handler:
            headers = {"Content-Type": "multipart/form-data"}
            user = ACTOR

            def require_user(self):
                return self.user

            def read_raw_body(self):
                return b"uploaded-body"

            def send_json(self, data, status):
                self.response = data, status

        handler = Handler()
        book = monthly_book(amount="invalid")
        with (
            patch.object(reimbursement_actions, "connect_db", return_value=self.conn),
            patch.object(
                reimbursement_actions, "parse_multipart_upload", return_value=("synthetic.xlsx", workbook_bytes(book))
            ),
        ):
            before = self.conn.total_changes
            reimbursement_actions.ReimbursementActionsMixin.handle_reimbursement_monthly_import(handler)
            self.assertEqual(handler.response[1], 400)
            self.assertIn("C2", handler.response[0]["error"])
            self.assertEqual(self.conn.total_changes, before)
        handler.user = {**ACTOR, "role": "room_admin"}
        with patch.object(reimbursement_actions, "connect_db") as connect:
            reimbursement_actions.ReimbursementActionsMixin.handle_reimbursement_monthly_import(handler)
            self.assertEqual(handler.response[1], 403)
            connect.assert_not_called()

    def test_import_reads_latest_manual_payment_after_acquiring_write_lock(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "synthetic.sqlite"
            writer = sqlite3.connect(path)
            writer.row_factory = sqlite3.Row
            self.addCleanup(writer.close)
            initialize_schema(writer)
            body = workbook_bytes(monthly_book())
            item = import_monthly_reimbursement_workbook(writer, body, ACTOR)["items"][0]
            writer.commit()
            writer.execute("BEGIN IMMEDIATE")
            latest = get_reimbursement_record(writer, item["id"])
            upsert_reimbursement_record(writer, {**latest, "paidAmount": 4.5})
            attempts = threading.Event()
            looked_up = threading.Event()
            outcomes = []
            original = application.get_reimbursement_record_by_key

            def lookup(*args):
                looked_up.set()
                return original(*args)

            def importing():
                conn = sqlite3.connect(path, timeout=3)
                conn.row_factory = sqlite3.Row
                conn.set_trace_callback(lambda sql: attempts.set() if sql.startswith("BEGIN IMMEDIATE") else None)
                try:
                    outcomes.append(import_monthly_reimbursement_workbook(conn, body, ACTOR)["items"][0])
                    conn.commit()
                except Exception as error:
                    outcomes.append(error)
                finally:
                    conn.close()

            with patch.object(application, "get_reimbursement_record_by_key", side_effect=lookup):
                thread = threading.Thread(target=importing)
                thread.start()
                try:
                    self.assertTrue(attempts.wait(2), "import must acquire the writer lock before business reads")
                    self.assertFalse(looked_up.is_set())
                finally:
                    writer.commit()
                    thread.join(4)
            self.assertFalse(thread.is_alive())
            self.assertEqual(len(outcomes), 1)
            self.assertIsInstance(outcomes[0], dict)
            self.assertEqual(outcomes[0]["paidAmount"], 4.5)
