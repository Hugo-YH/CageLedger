import copy
import json
import sqlite3
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import Mock, patch
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook

from server_app.domains.intake.strain_standard import standardize_strain, supplier_full_name
from server_app.domains.reception import repository, routes, service, workbook
from server_app.legacy import initialize_schema
from server_app.repositories.audit_storage import read_payload
from server_app.repositories.entities import upsert_intake_batch
from server_app.shared.concurrency import StaleWriteError
from server_app.shared.sqlite import ClosingConnection
from tests.test_intake_summary import reservation

ACTOR = {"id": "admin", "username": "admin", "displayName": "系统管理员", "role": "admin"}


class ReceptionTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        self.conn.row_factory = sqlite3.Row
        initialize_schema(self.conn)
        self.addCleanup(self.conn.close)
        self.index = {
            "source": "database",
            "updatedAt": "2026-10-31",
            "items": [
                {
                    "iacuc": "Z2026001",
                    "pi": "申请负责人；共同负责人",
                    "owner": "申请实验人",
                    "project": "申请项目",
                    "funding": "申请项目来源",
                    "fundCode": "0000123456789",
                    "projectEndDate": "2027-10-31",
                }
            ],
        }
        self.patcher = patch.object(repository, "read_iacuc_index", side_effect=lambda *_: self.index)
        self.patcher.start()
        self.addCleanup(self.patcher.stop)
        self.batch = reservation(
            status="received",
            quantity=20,
            sex="雄性/雌性",
            receipts=[{"actualReceiptDate": "2026-10-30", "cardCount": 4}],
            recognition={"sources": [{"kind": "husbandry", "maleQuantity": 12, "femaleQuantity": 8}]},
        )
        self.put(self.batch)

    def put(self, batch):
        upsert_intake_batch(self.conn, batch)
        self.conn.commit()

    def listed(self, **query):
        return service.list_records(self.conn, ACTOR, {key: [value] for key, value in query.items()})

    def body(self):
        result = self.listed()
        return {
            "ids": [item["id"] for item in result["items"]],
            "expectedVersions": {item["id"]: item["updatedAt"] for item in result["items"]},
            "expectedApplicationVersion": result["applicationVersion"],
        }

    def test_only_received_and_project_information_never_comes_from_reservation(self):
        self.put(reservation(2, status="printed"))
        before = self.conn.total_changes
        result = self.listed()
        self.assertEqual(result["page"]["total"], 1)
        item = result["items"][0]
        for field in ("pi", "owner", "project", "funding", "fundCode", "projectEndDate"):
            self.assertEqual(item[field], self.index["items"][0][field])
        self.assertEqual(item["receiptDates"], ["2026-10-30"])
        self.assertEqual(self.conn.total_changes, before)
        self.assertNotIn("rawMessage", item)
        self.assertNotIn("recognition", item)

    def test_missing_or_ambiguous_application_stays_blank(self):
        for source_items in ([], self.index["items"] * 2):
            with self.subTest(items=len(source_items)):
                self.index["items"] = source_items
                item = self.listed()["items"][0]
                self.assertFalse(item["projectMatched"])
                self.assertEqual(item["pi"], "")
                self.assertEqual(item["owner"], "")
                self.assertEqual(item["project"], "")
                self.assertTrue(any("未唯一匹配" in message for message in item["warnings"]))

    def test_authoritative_individual_people_filters_and_scoped_options(self):
        filters = json.dumps({"pi": ["共同负责人"]})
        self.assertEqual(self.listed(columnFilters=filters)["page"]["total"], 1)
        self.assertEqual(self.listed(columnFilters=json.dumps({"pi": ["预约人"]}))["page"]["total"], 0)
        values = service.options(self.conn, ACTOR, {"column": ["pi"], "columnFilters": [filters]})
        self.assertEqual({item["value"] for item in values["items"]}, {"申请负责人", "共同负责人"})
        self.assertEqual(self.listed(columnFilters=json.dumps({"batchNo": ["' OR 1=1 --"]}))["page"]["total"], 0)

    def test_empty_iacuc_cannot_match_an_empty_application(self):
        self.batch["iacuc"] = ""
        self.index["items"][0]["iacuc"] = ""
        self.put(self.batch)
        item = self.listed()["items"][0]
        self.assertFalse(item["projectMatched"])
        self.assertEqual(item["pi"], "")

    def test_pagination_dates_empty_filter_and_counts_match(self):
        self.put(reservation(2, status="received", receipts=[]))
        result = self.listed(limit="1")
        self.assertEqual((len(result["items"]), result["page"]["total"], result["page"]["hasMore"]), (1, 2, True))
        options = service.options(self.conn, ACTOR, {"column": ["receiptDate"]})
        self.assertEqual({item["value"] for item in options["items"]}, {"", "2026-10-30"})
        self.assertEqual(self.listed(columnFilters=json.dumps({"receiptDate": [""]}))["page"]["total"], 1)
        self.assertEqual(self.listed(columnFilters=json.dumps({"receiptDate": ["", "2026-10-30"]}))["page"]["total"], 2)

    def test_preview_splits_only_explicit_consistent_sex_quantities(self):
        document = service.preview(self.conn, ACTOR, self.body())
        self.assertEqual(document["rowCount"], 2)
        self.assertEqual([(line["sex"], line["quantity"]) for line in document["rows"]], [("雄性", 12), ("雌性", 8)])
        self.assertEqual(sum(line["quantity"] for line in document["rows"]), 20)
        self.assertEqual(document["rows"][0]["project"]["owner"], "申请实验人")
        self.batch["quantity"] = 30
        self.put(self.batch)
        document = service.preview(self.conn, ACTOR, self.body())
        self.assertEqual(document["rowCount"], 1)
        self.assertEqual(document["rows"][0]["quantity"], 30)
        self.assertTrue(any("无法唯一确认" in warning for warning in document["warnings"]))

    def test_multiple_receipts_do_not_duplicate_total_or_use_planned_date(self):
        for receipts in ([], [{"actualReceiptDate": "2026-10-29"}, {"actualReceiptDate": "2026-10-30"}]):
            with self.subTest(receipts=receipts):
                self.batch["receipts"] = receipts
                self.put(self.batch)
                document = service.preview(self.conn, ACTOR, self.body())
                self.assertTrue(all(line["actualReceiptDate"] == "" for line in document["rows"]))
                self.assertEqual(sum(line["quantity"] for line in document["rows"]), 20)

    def test_sex_evidence_from_another_batch_is_not_reused(self):
        self.batch["recognition"]["sources"][0]["batchNo"] = "different-batch"
        self.put(self.batch)
        document = service.preview(self.conn, ACTOR, self.body())
        self.assertEqual(document["rowCount"], 1)
        self.assertEqual(document["rows"][0]["quantity"], 20)

    def test_versions_and_status_are_checked_again_before_export(self):
        original = self.body()
        changed = copy.deepcopy(original)
        changed["expectedVersions"][self.batch["id"]] = "outdated"
        with self.assertRaises(StaleWriteError):
            service.preview(self.conn, ACTOR, changed)
        self.index["items"][0]["pi"] = "新的申请负责人"
        with self.assertRaises(StaleWriteError):
            service.preview(self.conn, ACTOR, original)
        refreshed = self.body()
        self.batch["status"] = "printed"
        self.put(self.batch)
        with self.assertRaises(StaleWriteError):
            service.preview(self.conn, ACTOR, refreshed)

    def test_export_preserves_original_B_to_U_and_text_ids_and_no_formulas(self):
        self.batch.update(purchaseOrderNo="000012345678901234567890", notes='=HYPERLINK("bad")')
        self.put(self.batch)
        document = service.preview(self.conn, ACTOR, self.body())
        book = load_workbook(BytesIO(workbook.generate(document)))
        sheet = book["接收登记"]
        self.assertEqual(tuple(sheet.cell(3, column).value or "" for column in range(2, 22)), workbook.HEADERS)
        self.assertEqual(sheet["B4"].value, self.batch["purchaseOrderNo"])
        self.assertEqual(sheet["P4"].value + sheet["P5"].value, 20)
        self.assertEqual(sheet["H4"].value.date().isoformat(), "2026-10-30")
        self.assertIsNone(sheet["M4"].value)
        self.assertEqual(sheet["K4"].value, "待补充")
        self.assertEqual(sheet["L4"].value, "SPF")
        self.assertEqual(sheet["L5"].value, "SPF")
        self.assertEqual(sheet["R4"].value, "待补充")
        self.assertEqual(sheet["U4"].data_type, "s")
        self.assertEqual(sheet.max_column, 21)
        self.assertFalse(any(cell.data_type == "f" for tab in book for row in tab for cell in row))
        book.close()

    def test_registration_grade_uses_species_without_rewriting_reservation(self):
        for species, expected in (
            ("mouse", "SPF"),
            ("rat", "SPF"),
            ("小鼠", "SPF"),
            ("大鼠", "SPF"),
            ("guinea_pig", "普通级"),
            ("rabbit", "普通级"),
            ("monkey", "普通级"),
            ("dog", "普通级"),
            ("pig", "普通级"),
            ("", ""),
        ):
            with self.subTest(species=species):
                self.batch["species"] = species
                self.put(self.batch)
                before = self.conn.execute("SELECT payload FROM intake_batches").fetchone()[0]
                document = service.preview(self.conn, ACTOR, self.body())
                self.assertTrue(document["rows"])
                self.assertTrue(all(line["grade"] == expected for line in document["rows"]))
                result = load_workbook(BytesIO(workbook.generate(document)))
                self.assertEqual(result["接收登记"]["L4"].value, expected or "待补充")
                self.assertEqual(self.conn.execute("SELECT payload FROM intake_batches").fetchone()[0], before)
                result.close()

    def test_export_uses_standardized_strain_and_full_supplier_without_rewriting_batch(self):
        self.batch.update(strainRaw="c57", strainStandard="C57BL/6J", supplier="江苏集萃")
        self.put(self.batch)
        before = self.conn.execute("SELECT payload FROM intake_batches").fetchone()[0]
        document = service.preview(self.conn, ACTOR, self.body())
        result = load_workbook(BytesIO(workbook.generate(document)))
        sheet = result["接收登记"]
        self.assertEqual(sheet["J4"].value, "C57BL/6J")
        self.assertEqual(sheet["Q4"].value, "江苏集萃药康生物科技股份有限公司")
        self.assertEqual(self.listed()["items"][0]["strain"], "C57BL/6J")
        self.assertEqual(self.listed()["items"][0]["supplier"], "江苏集萃")
        self.assertEqual(self.conn.execute("SELECT payload FROM intake_batches").fetchone()[0], before)
        result.close()

    def test_export_resolves_known_strain_alias_but_does_not_guess_unknown_strain(self):
        self.batch.update(strainRaw="c57", strainStandard="")
        self.put(self.batch)
        document = service.preview(self.conn, ACTOR, self.body())
        self.assertEqual(document["rows"][0]["strain"], standardize_strain("c57"))
        self.assertTrue(document["rows"][0]["strain"])
        self.assertEqual(self.listed()["items"][0]["strain"], document["rows"][0]["strain"])
        self.batch["strainRaw"] = "未收录的新动物品系"
        self.put(self.batch)
        document = service.preview(self.conn, ACTOR, self.body())
        self.assertEqual(document["rows"][0]["strain"], "")
        self.assertEqual(self.listed()["items"][0]["strain"], "")
        result = load_workbook(BytesIO(workbook.generate(document)))
        self.assertEqual(result["接收登记"]["J4"].value, "待补充")
        result.close()

    def test_full_supplier_preserves_branches_and_unknown_names(self):
        branch = "江苏集萃药康生物科技股份有限公司常州分公司"
        self.assertEqual(supplier_full_name(branch), (branch, True))
        self.assertEqual(supplier_full_name("广东药康"), ("广东药康生物科技有限公司", True))
        self.assertEqual(supplier_full_name("未知供应单位"), ("未知供应单位", False))
        self.assertEqual(supplier_full_name("江苏集萃 新增分公司"), ("江苏集萃 新增分公司", False))

    def test_authorization_and_input_boundaries(self):
        connection = Mock()
        with self.assertRaises(PermissionError):
            service.list_records(connection, None, {})
        connection.execute.assert_not_called()
        for role in ("admin", "room_admin", "viewer"):
            self.assertEqual(service.list_records(self.conn, {**ACTOR, "role": role}, {})["page"]["total"], 1)
        for query in (
            {"limit": ["101"]},
            {"offset": ["-1"]},
            {"sortKey": ["random"]},
            {"columnFilters": ['{"unknown":[]}']},
        ):
            with self.subTest(query=query), self.assertRaises(ValueError):
                service.list_records(self.conn, ACTOR, query)
        for body in ({}, {**self.body(), "ids": [self.batch["id"]] * 201}, {**self.body(), "expectedVersions": {}}):
            with self.subTest(body=body), self.assertRaises(ValueError):
                service.preview(self.conn, ACTOR, body)

    def test_excel_uses_original_registration_styles_and_clean_template(self):
        document = service.preview(self.conn, ACTOR, self.body())
        template = load_workbook(workbook.TEMPLATE)
        payload = workbook.generate(document)
        result = load_workbook(BytesIO(payload))
        sheet = result["接收登记"]
        reference = template["接收登记"]
        self.assertEqual(result.loaded_theme, template.loaded_theme)
        self.assertEqual(sheet.sheet_format.defaultRowHeight, 60)
        self.assertEqual(sheet.sheet_view.zoomScale, 40)
        self.assertEqual(sheet.sheet_view.zoomScaleNormal, 40)
        for row in (3, 4, 5):
            self.assertEqual(sheet.row_dimensions[row].height, 60)
        with ZipFile(BytesIO(payload)) as archive:
            styles = ElementTree.fromstring(archive.read("xl/styles.xml"))
        namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        normal = styles.find("s:cellStyles/s:cellStyle[@name='Normal']", namespace)
        normal_style = styles.find("s:cellStyleXfs", namespace)[int(normal.get("xfId"))]
        normal_font = styles.find("s:fonts", namespace)[int(normal_style.get("fontId"))]
        self.assertEqual(normal_font.find("s:name", namespace).get("val"), "等线")
        self.assertEqual(normal_font.find("s:sz", namespace).get("val"), "11")
        self.assertEqual(sheet["B3"].font.name, "宋体")
        self.assertEqual(sheet["B3"].font.sz, 20)
        self.assertTrue(sheet["B3"].font.bold)
        self.assertFalse(sheet["B4"].font.bold)
        self.assertEqual(sheet["D3"].fill.fgColor.rgb, "FFFFF3D1")
        self.assertEqual(sheet["Q4"].font.color.rgb, "FFFF0000")
        self.assertEqual(sheet["H4"].number_format, reference["H4"].number_format)
        for column in range(2, 22):
            letter = sheet.cell(3, column).column_letter
            self.assertEqual(sheet.column_dimensions[letter].width, reference.column_dimensions[letter].width)
            self.assertIsNone(reference.cell(4, column).value)
            for attribute in ("font", "fill", "border", "alignment"):
                self.assertEqual(
                    copy.copy(getattr(sheet.cell(3, column), attribute)),
                    copy.copy(getattr(reference.cell(3, column), attribute)),
                )
                for row in (4, 5):
                    if sheet.cell(row, column).value == "待补充" and attribute in ("font", "fill"):
                        self.assertEqual(sheet.cell(row, column).fill.fgColor.rgb, "FFFFC7CE")
                        self.assertEqual(sheet.cell(row, column).font.color.rgb, "FF9C0006")
                        self.assertEqual(sheet.cell(row, column).font.name, reference.cell(4, column).font.name)
                        self.assertEqual(sheet.cell(row, column).font.sz, reference.cell(4, column).font.sz)
                        continue
                    self.assertEqual(
                        copy.copy(getattr(sheet.cell(row, column), attribute)),
                        copy.copy(getattr(reference.cell(4, column), attribute)),
                    )
        allowed_values = {*workbook.HEADERS, "实验动物接收信息"}
        self.assertTrue(
            all(cell.value in allowed_values for row in reference for cell in row if cell.value is not None)
        )
        template.close()
        result.close()

    def test_missing_cells_are_marked_without_changing_zero_spacers_or_later_rows(self):
        document = service.preview(self.conn, ACTOR, self.body())
        first, second = document["rows"]
        first.update(supplier=" ", actualReceiptDate="invalid", quantity=0, notes="", grade="")
        first["project"]["pi"] = ""
        second["quantity"] = None
        book = load_workbook(BytesIO(workbook.generate(document)))
        template = load_workbook(workbook.TEMPLATE)
        sheet = book["接收登记"]
        for address in ("C4", "H4", "Q4", "P5", "K4", "L4", "R4", "S4"):
            self.assertEqual(sheet[address].value, "待补充")
            self.assertEqual(sheet[address].fill.fgColor.rgb, "FFFFC7CE")
            self.assertEqual(sheet[address].font.color.rgb, "FF9C0006")
        self.assertEqual(sheet["P4"].value, 0)
        self.assertIsNone(sheet["M4"].value)
        self.assertIsNone(sheet["U4"].value)
        self.assertEqual(sheet["C5"].value, "申请负责人；共同负责人")
        self.assertEqual(sheet["Q5"].value, second["supplier"])
        for column in ("C", "P", "Q"):
            self.assertEqual(
                copy.copy(sheet[f"{column}5" if column != "P" else "P4"].fill),
                copy.copy(template["接收登记"][f"{column}4"].fill),
            )
        template.close()
        book.close()

    def test_http_unauthenticated_never_opens_database(self):
        handler = Mock()
        handler.current_user.return_value = None
        with patch.object(routes, "connect_db") as connect:
            self.assertEqual(routes.read(handler, {}).status, 401)
            self.assertEqual(routes.prepare(handler, {}).status, 401)
            connect.assert_not_called()

    def test_http_export_logs_once_and_does_not_change_source_batches(self):
        body = self.body()
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "copy.sqlite"
            with sqlite3.connect(path, factory=ClosingConnection) as target:
                self.conn.backup(target)

            def connect():
                connection = sqlite3.connect(path, factory=ClosingConnection)
                connection.row_factory = sqlite3.Row
                return connection

            handler = Mock(path="/api/receptions/export")
            handler.current_user.return_value = ACTOR
            handler.read_json_body.return_value = body
            with patch.object(routes, "connect_db", side_effect=connect):
                self.assertTrue(routes.prepare(handler, {}))
            handler.send_download.assert_called_once()
            self.assertTrue(handler.send_download.call_args.args[0].startswith(b"PK"))
            with connect() as connection:
                audit_ids = connection.execute(
                    "SELECT id FROM audit_events WHERE action='reception.exported'"
                ).fetchall()
                self.assertEqual(len(audit_ids), 1)
                audit = json.loads(read_payload(connection, audit_ids[0]["id"]))
                self.assertEqual(audit["actorUserId"], ACTOR["id"])
                self.assertEqual(audit["after"]["rowCount"], 2)
                self.assertEqual(audit["after"]["batchIds"], body["ids"])
                self.assertEqual(
                    connection.execute("SELECT payload FROM intake_batches").fetchone()[0],
                    self.conn.execute("SELECT payload FROM intake_batches").fetchone()[0],
                )
