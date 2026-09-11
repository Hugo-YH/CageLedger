import copy
import json
import shutil
import sqlite3
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from server_app.domains.quarantine import files, service
from server_app.domains.quarantine import repository as repo
from server_app.domains.quarantine.documents import generate
from server_app.domains.quarantine.history import supplier_history
from server_app.shared.concurrency import StaleWriteError

ACTOR = {"id": "regular", "username": "regular", "displayName": "检测员", "role": "room_admin"}


def sample_test(batch_id="batch"):
    return {
        "id": "test",
        "batchId": batch_id,
        "method": "parasite",
        "samplingDate": "2026-09-02",
        "testDate": "2026-09-03",
        "conclusion": "所属检疫批次抽检无异常",
        "notes": "",
        "samples": [
            {
                "id": "sample",
                "number": "1",
                "material": "皮毛",
                "poolCount": 1,
                "portionCount": 2,
                "sourceIds": ["source-a"],
            }
        ],
        "projects": [
            {
                "id": "endo",
                "name": "体内寄生虫",
                "kit": "",
                "lot": "",
                "sampleIds": ["sample"],
                "results": {"sample": "negative"},
                "wells": {},
                "nc": "",
                "pc": "",
            }
        ],
    }


class QuarantineTests(unittest.TestCase):
    def setUp(self):
        self.conn = sqlite3.connect(":memory:")
        self.conn.row_factory = sqlite3.Row
        self.conn.execute("PRAGMA foreign_keys=ON")
        repo.ensure_schema(self.conn)
        self.conn.execute("CREATE TABLE intake_batches(id TEXT PRIMARY KEY, intake_date TEXT, payload TEXT)")
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.audits = []
        self.audit_patch = patch(
            "server_app.domains.quarantine.service.write_audit_events",
            side_effect=lambda conn, events: self.audits.extend(events),
        )
        self.audit_patch.start()
        sources = [
            {"id": "source-a", "supplier": "供应商甲", "species": "小鼠"},
            {"id": "source-b", "supplier": "供应商乙", "species": "小鼠"},
        ]
        self.batch = self.write(service.save_batch, {"item": {"id": "batch", "name": "周三检疫", "sources": sources}})

    def tearDown(self):
        self.audit_patch.stop()
        self.conn.close()
        self.tmp.cleanup()

    def write(self, operation, *args, **kwargs):
        with self.conn:
            return operation(self.conn, ACTOR, *args, **kwargs)

    def create_test(self, raw=None):
        return self.write(service.save_test, {"item": raw or sample_test()})

    def issue(self, test):
        return self.write(
            files.issue,
            test["id"],
            {"expectedUpdatedAt": test["updatedAt"], "expectedBatchUpdatedAt": self.batch["updatedAt"]},
            self.root,
        )

    def test_schema_is_idempotent_and_coverage_does_not_require_sampling(self):
        repo.ensure_schema(self.conn)
        test = self.create_test()
        self.assertEqual(len(service.batch_detail(self.conn, "batch")["item"]["sources"]), 2)
        self.assertEqual(test["samples"][0]["sourceIds"], ["source-a"])
        self.assertEqual(self.audits[-1]["actorUserId"], "regular")

    def test_missing_and_stale_versions_rejected(self):
        test = self.create_test()
        for expected in ("", "old"):
            with self.assertRaises(StaleWriteError):
                self.write(service.save_test, {"item": test, "expectedUpdatedAt": expected}, test["id"])
        updated = self.write(service.save_test, {"item": test, "expectedUpdatedAt": test["updatedAt"]}, test["id"])
        self.assertNotEqual(test["updatedAt"], updated["updatedAt"])

    def test_source_snapshot_is_not_overwritten_by_intake_edit(self):
        with self.conn:
            self.conn.execute(
                "INSERT INTO intake_batches VALUES (?, ?, ?)",
                (
                    "intake",
                    "2026-09-02",
                    json.dumps(
                        {"supplier": "原供应商", "species": "小鼠", "intakeDate": "2026-09-02", "status": "received"}
                    ),
                ),
            )
        batch = self.write(
            service.save_batch,
            {"item": {"id": "snapshot", "name": "快照", "sources": [{"id": "snap", "intakeId": "intake"}]}},
        )
        with self.conn:
            self.conn.execute(
                "UPDATE intake_batches SET payload=?", (json.dumps({"supplier": "新供应商", "species": "大鼠"}),)
            )
        batch = self.write(service.save_batch, {"item": batch, "expectedUpdatedAt": batch["updatedAt"]}, batch["id"])
        self.assertEqual(batch["sources"][0]["supplier"], "原供应商")

    def test_issued_report_immutable_idempotent_and_corrections_preserved(self):
        test = self.create_test()
        report = self.issue(test)
        original = (self.root / report["storageName"]).read_bytes()
        self.assertEqual(self.issue(test)["id"], report["id"])
        current = repo.get(self.conn, "tests", test["id"])
        with self.assertRaises(ValueError):
            self.write(service.save_test, {"item": current, "expectedUpdatedAt": current["updatedAt"]}, current["id"])
        corrected = self.write(
            service.clone_test, test["id"], {"id": "correction", "expectedUpdatedAt": current["updatedAt"]}
        )
        report2 = self.issue(corrected)
        self.assertEqual(report2["version"], 2)
        self.assertEqual((self.root / report["storageName"]).read_bytes(), original)
        self.assertEqual(len(repo.all_items(self.conn, "reports")), 2)

    def test_generation_failure_rolls_back_and_retry_succeeds(self):
        test = self.create_test()
        with (
            patch("server_app.domains.quarantine.files.generate", side_effect=OSError("disk full")),
            self.assertRaises(OSError),
        ):
            self.issue(test)
        self.assertEqual(repo.get(self.conn, "tests", test["id"])["state"], "draft")
        self.assertEqual(repo.all_items(self.conn, "reports"), [])
        self.issue(test)

    def test_blank_is_not_negative_and_positive_can_issue(self):
        raw = sample_test()
        raw["projects"][0]["results"] = {}
        test = self.create_test(raw)
        with self.assertRaises(ValueError):
            self.issue(test)
        test["projects"][0]["results"]["sample"] = "positive"
        test["conclusion"] = "抽检发现异常，待处理"
        test = self.write(service.save_test, {"item": test, "expectedUpdatedAt": test["updatedAt"]}, test["id"])
        self.issue(test)
        self.assertEqual(service.batch_detail(self.conn, "batch")["item"]["status"], "存在异常，待处理")

    def test_cross_supplier_retests_do_not_rewrite_initial_abnormality(self):
        raw = sample_test()
        raw["samples"][0]["sourceIds"].append("source-b")
        raw["projects"][0]["results"]["sample"] = "suspect"
        test = self.create_test(raw)
        rows = supplier_history(self.conn, {})
        self.assertTrue(all(r["pending"] == 1 and r["confirmed"] == 0 for r in rows))
        retest = self.write(
            service.clone_test,
            test["id"],
            {"id": "retest", "supplier": "供应商甲", "expectedUpdatedAt": test["updatedAt"]},
            retest=True,
        )
        self.assertEqual(retest["samples"][0]["sourceIds"], ["source-a"])
        self.assertEqual(retest["projects"][0]["results"], {})
        self.assertEqual(service.batch_detail(self.conn, "batch")["item"]["status"], "待分开复检")
        self.assertEqual(repo.get(self.conn, "tests", test["id"])["projects"][0]["results"]["sample"], "suspect")

    def test_elisa_species_and_sample_applicability(self):
        raw = sample_test()
        raw["method"] = "elisa_rat"
        with self.assertRaises(ValueError):
            self.create_test(raw)
        raw["method"] = "pcr"
        raw["projects"][0]["sampleIds"] = ["missing"]
        with self.assertRaises(ValueError):
            self.create_test(raw)

    def test_upload_provenance_and_issued_rejection(self):
        test = self.create_test()
        uploaded = self.write(
            files.upload,
            test["id"],
            {"expectedUpdatedAt": test["updatedAt"], "sampleId": "sample"},
            "原始记录.pdf",
            b"%PDF-1.7\n",
            self.root,
        )
        self.assertEqual(uploaded["item"]["uploadedBy"]["id"], ACTOR["id"])
        self.issue(uploaded["test"])
        current = repo.get(self.conn, "tests", test["id"])
        with self.assertRaises(ValueError):
            self.write(
                files.upload, test["id"], {"expectedUpdatedAt": current["updatedAt"]}, "new.pdf", b"%PDF-1.7", self.root
            )
        self.assertEqual(len(repo.all_items(self.conn, "attachments")), 1)

    def test_word_has_draft_marker_repeated_headers_and_blank_signatures(self):
        test = self.create_test()
        content = generate(files.snapshot(self.conn, test), self.root, draft=True)
        with ZipFile(BytesIO(content)) as archive:
            xml = "".join(archive.read(name).decode() for name in archive.namelist() if name.endswith(".xml"))
        for token in ("草稿", "检测人：", "复核人：", "tblHeader", "体内寄生虫"):
            self.assertIn(token, xml)
        self.assertNotIn("检测员", xml)

    def test_word_preserves_source_styles_page_system_and_horizontal_tables(self):
        from docx import Document

        from server_app.domains.quarantine.documents import TEMPLATES

        test = self.create_test()
        content = generate(files.snapshot(self.conn, test), self.root, draft=True)
        original = Document(TEMPLATES / "parasite.docx")
        exported = Document(BytesIO(content))
        self.assertEqual(exported.sections[0]._sectPr.xml, original.sections[0]._sectPr.xml)
        self.assertEqual(exported.tables[0].cell(0, 0).text, original.tables[0].cell(0, 0).text)
        self.assertEqual(exported.tables[1].cell(0, 0).text, "样本编号")
        self.assertEqual(exported.tables[1].cell(1, 0).text, "动物批号")
        with ZipFile(BytesIO(content)) as output, ZipFile(TEMPLATES / "parasite.docx") as template:
            self.assertEqual(output.read("word/styles.xml"), template.read("word/styles.xml"))

    def test_report_form_derives_counts_and_issues_without_extra_conclusion(self):
        raw = sample_test()
        raw.update(reportFormVersion=2, reportMaterial="血清", reportSpecimenState="液体", conclusion="")
        raw["samples"][0].update(sourceIds=["source-a", "source-b", "source-a"], poolCount=99, portionCount=99)
        record = self.create_test(raw)
        self.assertEqual(record["samples"][0]["sourceIds"], ["source-a", "source-b"])
        self.assertEqual(record["samples"][0]["poolCount"], 1)
        self.assertEqual(record["samples"][0]["portionCount"], 2)
        with ZipFile(BytesIO(generate(files.snapshot(self.conn, record), self.root, draft=True))) as archive:
            xml = archive.read("word/document.xml").decode()
        self.assertIn("1样（2份血清）", xml)
        self.assertNotIn("结果判定：", xml)
        self.assertIn("代表空白", xml)
        self.issue(record)

    def test_multiple_project_metadata_versions_and_original_provenance(self):
        from server_app.domains.quarantine import attachment_metadata

        raw = sample_test()
        raw["projects"].append({**copy.deepcopy(raw["projects"][0]), "id": "ecto", "name": "体外寄生虫"})
        record = self.create_test(raw)
        uploaded = self.write(
            files.upload,
            record["id"],
            {"expectedUpdatedAt": record["updatedAt"], "projectIds": '["endo","ecto"]'},
            "raw.pdf",
            b"%PDF-1.7",
            self.root,
        )
        original = uploaded["item"]
        self.assertEqual(original["projectIds"], ["endo", "ecto"])
        updated = self.write(
            attachment_metadata.update,
            original["id"],
            {
                "expectedUpdatedAt": uploaded["test"]["updatedAt"],
                "item": {"caption": "两个项目的原始记录", "position": 2},
            },
        )
        self.assertEqual(updated["item"]["uploadedAt"], original["uploadedAt"])
        self.assertEqual(updated["item"]["uploadedBy"], original["uploadedBy"])
        self.assertEqual(updated["item"]["position"], 2)
        with self.assertRaises(StaleWriteError):
            self.write(
                attachment_metadata.update,
                original["id"],
                {"expectedUpdatedAt": uploaded["test"]["updatedAt"], "item": {"removed": True}},
            )
        self.issue(updated["test"])
        issued = repo.get(self.conn, "tests", record["id"])
        with self.assertRaises(ValueError):
            self.write(
                attachment_metadata.update,
                original["id"],
                {"expectedUpdatedAt": issued["updatedAt"], "item": {"removed": True}},
            )
        correction = self.write(
            service.clone_test, record["id"], {"id": "corrected", "expectedUpdatedAt": issued["updatedAt"]}
        )
        copied = repo.all_items(self.conn, "attachments", correction["id"])[0]
        self.write(
            attachment_metadata.update,
            copied["id"],
            {"expectedUpdatedAt": correction["updatedAt"], "item": {"removed": True}},
        )
        self.assertEqual(repo.all_items(self.conn, "attachments", correction["id"]), [])
        self.assertEqual(len(repo.all_items(self.conn, "attachments", record["id"])), 1)

    def test_anonymous_domain_write_rejected(self):
        with self.assertRaises(PermissionError):
            service.save_batch(self.conn, None, {"item": {}})

    def test_supplier_counts_deduplicate_intakes_and_separate_methods(self):
        first = self.create_test()
        second = copy.deepcopy(sample_test())
        second.update(id="pcr", method="pcr")
        self.create_test(second)
        rows = supplier_history(self.conn, {"supplier": "供应商甲"})
        self.assertEqual(len(rows), 2)
        self.assertTrue(all(r["poolCount"] == 1 and r["portionCount"] == 2 and r["batchCount"] == 1 for r in rows))
        self.issue(first)
        current = repo.get(self.conn, "tests", first["id"])
        correction = self.write(
            service.clone_test, first["id"], {"id": "correct", "expectedUpdatedAt": current["updatedAt"]}
        )
        self.issue(correction)
        self.assertEqual(sum(r["poolCount"] for r in supplier_history(self.conn, {"supplier": "供应商甲"})), 2)

    def test_split_retest_completion_and_manual_conclusion(self):
        raw = sample_test()
        raw["samples"][0]["sourceIds"] = ["source-a", "source-b"]
        raw["projects"][0]["results"]["sample"] = "suspect"
        original = self.create_test(raw)
        for index, supplier in enumerate(("供应商甲", "供应商乙")):
            retest = self.write(
                service.clone_test,
                original["id"],
                {"id": f"retest-{index}", "supplier": supplier, "expectedUpdatedAt": original["updatedAt"]},
                retest=True,
            )
            retest.update(samplingDate="2026-09-04", testDate="2026-09-04", conclusion="分开复检阴性")
            retest["projects"][0]["results"]["sample"] = "negative"
            retest = self.write(
                service.save_test, {"item": retest, "expectedUpdatedAt": retest["updatedAt"]}, retest["id"]
            )
            self.issue(retest)
        self.assertEqual(service.batch_detail(self.conn, "batch")["item"]["status"], "复检完成，待确认结论")
        initial_rows = [r for r in supplier_history(self.conn, {}) if not r["retest"]]
        self.assertTrue(all(r["resolved"] == 1 and r["pending"] == 0 for r in initial_rows))
        self.batch["conclusion"] = "分开复检无异常"
        self.write(service.save_batch, {"item": self.batch, "expectedUpdatedAt": self.batch["updatedAt"]}, "batch")
        self.assertIn("保留异常记录", service.batch_detail(self.conn, "batch")["item"]["status"])

    def test_pcr_rat_only_projects_reject_mouse_and_allow_partial_samples(self):
        raw = sample_test()
        raw["method"] = "pcr"
        raw["projects"][0]["name"] = "支气管鲍特杆菌（大鼠）"
        with self.assertRaises(ValueError):
            self.create_test(raw)
        raw["projects"][0]["sampleIds"] = []
        raw["projects"][0]["results"] = {}
        test = self.create_test(raw)
        self.assertEqual(test["projects"][0]["sampleIds"], [])


class QuarantineBackupTests(unittest.TestCase):
    def test_database_and_relative_files_restore_together(self):
        with tempfile.TemporaryDirectory() as temporary:
            original = Path(temporary) / "original"
            original.mkdir()
            conn = sqlite3.connect(original / "quarantine.sqlite")
            conn.row_factory = sqlite3.Row
            repo.ensure_schema(conn)
            with patch("server_app.domains.quarantine.service.write_audit_events"):
                with conn:
                    batch = service.save_batch(
                        conn,
                        ACTOR,
                        {
                            "item": {
                                "id": "batch",
                                "name": "备份检疫",
                                "sources": [{"id": "source-a", "supplier": "供应商", "species": "小鼠"}],
                            }
                        },
                    )
                with conn:
                    test = service.save_test(conn, ACTOR, {"item": sample_test()})
                with conn:
                    report = files.issue(
                        conn,
                        ACTOR,
                        test["id"],
                        {"expectedUpdatedAt": test["updatedAt"], "expectedBatchUpdatedAt": batch["updatedAt"]},
                        original / "files" / "quarantine",
                    )
            conn.close()
            restored = Path(temporary) / "restored"
            shutil.copytree(original, restored)
            with sqlite3.connect(restored / "quarantine.sqlite") as restored_conn:
                restored_report = repo.get(restored_conn, "reports", report["id"])
            restored_conn.close()
            self.assertEqual(
                (restored / "files" / "quarantine" / restored_report["storageName"]).read_bytes(),
                (original / "files" / "quarantine" / report["storageName"]).read_bytes(),
            )
