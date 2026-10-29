import io
import sqlite3
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

import server
from server_app.domains.billing import monthly_summary, settlement_xlsx_layout
from server_app.domains.billing.settlement_export import build_settlement_workbook
from server_app.domains.quantity.application import save_quantity_sheet
from server_app.domains.technical_services import xlsx as technical_xlsx
from server_app.legacy import initialize_schema
from tests.test_custom_billing import AllQuantityCustomBillingTests
from tests.test_settlement_xlsx import sample_entry

MARKER = '=HYPERLINK("sample")'
ACTOR = {"id": "admin", "username": "admin", "displayName": "Admin", "role": "admin", "roomIds": []}


def workbook(content):
    return load_workbook(io.BytesIO(content), data_only=False)


def monthly_row(value):
    fields = (
        "pi",
        "facility",
        "iacuc",
        "funding",
        "species",
        "statementReturned",
        "fundBookNo",
        "reimbursementFormNo",
        "reimbursementNotes",
        "projectStartDate",
        "projectEndDate",
        "notes",
        "roomNames",
    )
    return {**dict.fromkeys(fields, value), "amount": 10, "supportAmount": 2, "reimbursementAmount": 3}


def technical_statement(value):
    return {
        "sourceType": "pi_merged_technical_service",
        "pi": value,
        "month": "2026-07",
        "totalAmount": 1,
        "technicalServiceRows": [
            {
                "iacuc": value,
                "facility": "zhujiang",
                "serviceName": value,
                "date": "2026-07-01",
                "unit": value,
                "quantity": 1,
                "unitPrice": 1,
                "amount": 1,
            }
        ],
    }


class XlsxLiteralTextTests(unittest.TestCase):
    def test_saved_quantity_fields_and_custom_notes_remain_literal(self):
        sheet, room = AllQuantityCustomBillingTests().fixture(False)
        sheet.update(pi=MARKER, owner=MARKER, funding=MARKER)
        sheet["customBillingSegments"][0]["note"] = MARKER
        conn = sqlite3.connect(":memory:")
        conn.row_factory = sqlite3.Row
        try:
            initialize_schema(conn)
            saved, *_ = save_quantity_sheet(conn, sheet, ACTOR)
            lines = server.quantity_sheet_statement_lines([saved], 0, [room], {})
            statement = {
                "id": "s",
                "iacucs": [saved["iacuc"]],
                "month": saved["month"],
                "pi": saved["pi"],
                "owner": saved["owner"],
                "funding": saved["funding"],
                "billingUnit": "cage_day",
            }
            result = workbook(build_settlement_workbook([(statement, lines)])).active
            self.assertEqual(
                (result["A1"].value, result["A1"].data_type), (MARKER + "课题组实验动物饲养费核算汇总表", "s")
            )
            for address in ("E2", "E3", "B5"):
                self.assertEqual((result[address].value, result[address].data_type), (MARKER, "s"))
            notes = [cell for row in result for cell in row if cell.value == MARKER and cell.column == 7]
            self.assertTrue(notes)
            self.assertTrue(all(cell.data_type == "s" for cell in notes))
            self.assertTrue(
                any(cell.data_type == "f" and cell.value.startswith("=SUM(") for row in result for cell in row)
            )
        finally:
            conn.rollback()
            conn.close()

    def test_monthly_and_technical_data_are_text_and_formulas_remain_formulas(self):
        monthly = workbook(monthly_summary.build_monthly_summary_xlsx("2026-07", [monthly_row(MARKER)])).active
        for column in "BCDEFJKL NOPQR".replace(" ", ""):
            cell = monthly[f"{column}3"]
            self.assertEqual((cell.value, cell.data_type), (MARKER, "s"))
        self.assertEqual((monthly["I3"].value, monthly["I3"].data_type), ("=G3-H3", "f"))
        self.assertEqual((monthly["G3"].value, monthly["G3"].data_type), (10, "n"))
        technical = workbook(build_settlement_workbook([(technical_statement(MARKER), [])])).active
        for address in ("D2", "A5", "C5", "E5"):
            self.assertEqual((technical[address].value, technical[address].data_type), (MARKER, "s"))
        self.assertEqual((technical["H6"].value, technical["H6"].data_type), ("=SUM(H5:H5)", "f"))
        self.assertEqual((technical["H5"].value, technical["H5"].data_type), (1, "n"))

    def test_ordinary_workbooks_preserve_values_types_styles_and_layout(self):
        fixtures = [
            lambda: build_settlement_workbook([sample_entry(include_custom=True)]),
            lambda: monthly_summary.build_monthly_summary_xlsx("2026-07", [monthly_row("普通文本")]),
            lambda: build_settlement_workbook([(technical_statement("普通文本"), [])]),
        ]

        def original_write(cell, value):
            cell.value = value
            return cell

        for fixture in fixtures:
            with self.subTest(fixture=fixture):
                actual = workbook(fixture()).active
                with (
                    patch.object(settlement_xlsx_layout, "set_literal_cell", original_write),
                    patch.object(monthly_summary, "set_literal_cell", original_write),
                    patch.object(technical_xlsx, "set_literal_cell", original_write),
                ):
                    before = workbook(fixture()).active
                self.assertEqual(str(actual.merged_cells), str(before.merged_cells))
                self.assertEqual(actual.freeze_panes, before.freeze_panes)
                self.assertEqual((actual.max_row, actual.max_column), (before.max_row, before.max_column))
                for row in actual:
                    for cell in row:
                        previous = before[cell.coordinate]
                        self.assertEqual(
                            (cell.value, cell.data_type, cell._style),
                            (previous.value, previous.data_type, previous._style),
                        )


if __name__ == "__main__":
    unittest.main()
