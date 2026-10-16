import io
import json
import sqlite3
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from server_app.cache import invalidate_data_cache, invalidate_data_cache_prefixes
from server_app.domains.billing.monthly_summary import build_monthly_summary_xlsx, build_pi_reimbursement_context
from server_app.domains.reimbursement_ledger.service import (
    confirm_allocation,
    create_allocation,
    list_obligations,
    save_claim,
)
from server_app.domains.technical_services import repository as repo
from server_app.domains.technical_services import service, settlement
from server_app.domains.technical_services.rules import daily_rows
from server_app.domains.technical_services.summary import merge_service_rows
from server_app.domains.workflow.application import save_billing_statement_workflow, update_workflow_status
from server_app.domains.workflow.permissions import require_workflow_withdrawal, withdrawal_capabilities
from server_app.pdf.documents import billing_statement_html
from server_app.persistence.bootstrap import initialize_schema
from server_app.shared.concurrency import StaleWriteError

ADMIN = {"id": "admin", "displayName": "管理员", "role": "admin", "username": "admin"}
STAFF = {"id": "staff", "displayName": "同名管理员", "role": "room_admin", "username": "staff"}


class TechnicalServiceTests(unittest.TestCase):
    def setUp(self):
        invalidate_data_cache()
        invalidate_data_cache_prefixes("applications_by_iacuc")
        self.conn = sqlite3.connect(":memory:")
        self.conn.row_factory = sqlite3.Row
        self.conn.execute("PRAGMA foreign_keys=ON")
        with patch("server_app.persistence.bootstrap.ensure_seed_images"):
            initialize_schema(self.conn)
        application = {"id": "app", "iacuc": "Z2026001", "pi": "纪老师", "project": "眼科项目", "funding": "项目来源"}
        self.conn.execute(
            "INSERT INTO experiment_applications(id,iacuc,imported_at,payload) VALUES(?,?,?,?)",
            ("app", "Z2026001", "2026-10-01", json.dumps(application)),
        )
        self.conn.commit()
        self.addCleanup(self.conn.close)
        self.addCleanup(invalidate_data_cache)
        self.addCleanup(invalidate_data_cache_prefixes, "applications_by_iacuc")

    def sheet(self, **fields):
        return {
            "id": "service-1",
            "month": "2026-10",
            "iacuc": "Z2026001",
            "facility": "zhujiang",
            "serviceId": "large_eye_drops",
            "customerType": "internal",
            "rows": [
                {"date": f"2026-10-{day:02}", "quantity": quantity}
                for day, quantity in ((1, 2), (2, 2), (3, 4), (10, 2))
            ],
            **fields,
        }

    def save(self, source=None, actor=ADMIN, identity=None, expected=None):
        with self.conn:
            return service.save(
                self.conn, actor, {"sheet": source or self.sheet(), "expectedUpdatedAt": expected}, identity
            )["item"]

    def start(self):
        preview = settlement.preview(self.conn, ADMIN, {"month": "2026-10", "pi": "纪老师"})
        with self.conn:
            return settlement.initiate(
                self.conn, ADMIN, {"month": "2026-10", "pi": "纪老师", "sourceVersions": preview["sourceVersions"]}
            )

    def test_sheet_column_filters_facets_and_pagination(self):
        for index, item in enumerate(repo.catalog(self.conn)[:12]):
            self.save(self.sheet(id=f"sheet-{index}", serviceId=item["id"]))
        first = repo.list_sheets(self.conn, {"limit": 5})
        self.assertEqual(first["page"]["total"], 12)
        target = repo.catalog(self.conn)[11]["name"]
        filters = {"columnFilters": json.dumps({"serviceName": [target]}), "limit": 5}
        selected = repo.list_sheets(self.conn, filters)
        self.assertEqual(selected["page"]["total"], 1)
        self.assertEqual(selected["items"][0]["serviceName"], target)
        facets = repo.sheet_filter_options(self.conn, {**filters, "column": "serviceName"})["items"]
        self.assertEqual(sum(option["count"] for option in facets), 12)
        self.assertEqual(len(facets), 12)
        zero = repo.list_sheets(self.conn, {"columnFilters": {"facility": ["bioisland"]}})
        self.assertEqual(zero["page"]["total"], 0)
        sorted_items = repo.list_sheets(self.conn, {"sortKey": "amount", "sortDir": "asc"})["items"]
        self.assertEqual(
            [item["totalAmount"] for item in sorted_items], sorted(item["totalAmount"] for item in sorted_items)
        )
        with self.assertRaises(ValueError):
            repo.list_sheets(self.conn, {"sortKey": "month; DROP TABLE technical_service_sheets"})
        with self.assertRaises(ValueError):
            repo.sheet_filter_options(self.conn, {"column": "month", "columnFilters": '{"month":"bad"}'})

    def test_candidate_column_filters_keep_full_pi_scope_and_frozen_amount(self):
        self.save()
        application = {"id": "app-2", "iacuc": "Z2026002", "pi": "纪老师", "project": "第二项目", "funding": "经费"}
        self.conn.execute(
            "INSERT INTO experiment_applications(id,iacuc,imported_at,payload) VALUES(?,?,?,?)",
            ("app-2", "Z2026002", "2026-10-01", json.dumps(application)),
        )
        self.conn.commit()
        invalidate_data_cache_prefixes("applications_by_iacuc")
        self.save(self.sheet(id="second", iacuc="Z2026002", facility="bioisland"))
        filtered = settlement.candidates(self.conn, ADMIN, {"columnFilters": {"iacuc": ["Z2026002"]}})
        self.assertEqual(filtered["items"][0]["totalAmount"], 600)
        self.assertEqual(filtered["items"][0]["iacucs"], ["Z2026001", "Z2026002"])
        facets = repo.candidate_filter_options(self.conn, {"column": "iacuc", "columnFilters": {"iacuc": ["Z2026002"]}})
        self.assertEqual({item["value"] for item in facets["items"]}, {"Z2026001", "Z2026002"})
        self.start()
        self.save(self.sheet(id="extra", serviceId="veterinary_observation"))
        frozen_filters = {"columnFilters": {"amount": ["600.0"], "workflow": ["statement_sent"]}}
        scopes = settlement.candidates(self.conn, ADMIN, frozen_filters)
        self.assertEqual(scopes["page"]["total"], 1)
        self.assertEqual(scopes["items"][0]["totalAmount"], 600)
        self.assertEqual(
            repo.sheet_filter_options(self.conn, {"column": "status"})["items"],
            [
                {"value": "editable", "label": "可编辑", "count": 1},
                {"value": "protected", "label": "已发起结算", "count": 2},
            ],
        )
        self.assertEqual(
            repo.candidate_filter_options(self.conn, {"column": "amount"})["items"],
            [{"value": "600.0", "label": "600.00", "count": 1}],
        )

    def test_unified_candidates_keep_fee_identity_filters_and_pagination(self):
        from server_app.domains.billing.unified_candidates import filter_options, list_candidates
        from tests.test_billing_candidates import seed_quantity_sheet

        seed_quantity_sheet(self.conn, "feeding-1", "2026-10", "纪老师", "Z2026001")
        self.conn.commit()
        self.save()
        filters = {"limit": 10, "offset": 0, "sortKey": "month", "sortDir": "desc", "columnFilters": {}}

        def calculate(month, pi):
            return {"totalAmount": 55}

        result = list_candidates(self.conn, ADMIN, filters, calculate)
        self.assertEqual(result["page"]["total"], 2)
        self.assertEqual(len({item["id"] for item in result["items"]}), 2)
        self.assertEqual(
            {(item["feeType"], item["totalAmount"]) for item in result["items"]},
            {("feeding", 55), ("technical_service", 300)},
        )
        self.assertEqual(
            list_candidates(self.conn, ADMIN, {**filters, "limit": 1, "offset": 1}, calculate)["page"]["total"], 2
        )
        selected = {**filters, "columnFilters": {"feeType": ["technical_service"], "amount": ["300.00"]}}
        self.assertEqual(
            list_candidates(self.conn, ADMIN, selected, calculate)["items"][0]["sourceType"], repo.SOURCE_TYPE
        )
        options = filter_options(
            self.conn, ADMIN, {**filters, "columnFilters": {"feeType": ["technical_service"]}}, "feeType", calculate
        )["items"]
        self.assertEqual({item["value"] for item in options}, {"feeding", "technical_service"})
        self.conn.commit()
        self.start()
        self.save(self.sheet(id="new-service", serviceId="veterinary_observation"))
        self.assertEqual(list_candidates(self.conn, ADMIN, selected, calculate)["items"][0]["totalAmount"], 300)
        reader = {**ADMIN, "role": "viewer"}
        self.assertEqual(list_candidates(self.conn, reader, filters, calculate)["page"]["total"], 1)
        with self.assertRaises(ValueError):
            list_candidates(self.conn, ADMIN, {**filters, "sortKey": "bad"}, calculate)

    def test_service_statement_excel_uses_confirmed_snapshot_and_units(self):
        from server_app.domains.billing.settlement_export import build_settlement_workbook

        self.save()
        preview = settlement.preview(self.conn, ADMIN, {"month": "2026-10", "pi": "纪老师"})
        workbook = load_workbook(io.BytesIO(build_settlement_workbook([(preview["statement"], preview["lines"])])))
        sheet = workbook.active
        self.assertEqual(sheet["B3"].value, 4)
        self.assertEqual(sheet["D3"].value, 300)
        self.assertEqual(sheet["E5"].value, "眼")
        self.assertEqual(sum(sheet[f"F{row}"].value for row in range(5, 9)), 10)
        self.assertEqual(sheet["H9"].value, "=SUM(H5:H8)")
        self.assertEqual(sheet["D3"].number_format, "0.00")

    def test_sample_catalog_migration_and_replay(self):
        repo.ensure_schema(self.conn)
        self.conn.commit()
        self.assertEqual(len(repo.catalog(self.conn)), 21)
        item = self.save()
        self.assertEqual((item["serviceDayCount"], item["totalQuantity"], item["totalAmount"]), (4, 10, 300))
        replay = self.save()
        self.assertEqual(replay["updatedAt"], item["updatedAt"])
        self.assertEqual(
            self.conn.execute("SELECT COUNT(*) FROM audit_events WHERE action='technical_service.created'").fetchone()[
                0
            ],
            1,
        )
        self.assertEqual(self.conn.execute("SELECT COUNT(*) FROM technical_service_rows").fetchone()[0], 4)
        self.assertEqual(item["funding"], "项目来源")

    def test_scope_duplicate_conflict_and_version(self):
        item = self.save()
        with self.assertRaises(StaleWriteError):
            self.save(self.sheet(id="different"))
        with self.assertRaises(StaleWriteError):
            self.save(self.sheet(), identity=item["id"], expected="outdated")
        with self.assertRaises(StaleWriteError):
            self.save(self.sheet(), identity=item["id"])
        revised = self.sheet(rows=[{"date": "2026-10-01", "quantity": 1}])
        self.assertEqual(self.save(revised, identity=item["id"], expected=item["updatedAt"])["totalAmount"], 30)

    def test_unsafe_or_non_string_draft_ids_are_rejected_before_writing(self):
        for identity in ("../", "a/b", True):
            with self.subTest(identity=identity), self.assertRaises(ValueError):
                self.save(self.sheet(id=identity))
        self.assertEqual(repo.list_sheets(self.conn, {})["page"]["total"], 0)

    def test_delete_tombstone_cannot_be_replayed(self):
        item = self.save()
        with self.conn:
            service.delete(self.conn, ADMIN, item["id"], item["updatedAt"])
        with self.assertRaises(StaleWriteError):
            self.save()
        self.assertEqual(repo.list_sheets(self.conn, {})["page"]["total"], 0)

    def test_units_dates_and_permissions(self):
        for value in (True, -1, float("inf"), 1.5):
            with self.assertRaises(ValueError):
                daily_rows([{"date": "2026-10-01", "quantity": value}], "2026-10", "眼", 30)
        with self.assertRaises(ValueError):
            daily_rows([{"date": "2026-11-01", "quantity": 1}], "2026-10", "眼", 30)
        with self.assertRaises(ValueError):
            daily_rows([{"date": "2026-10-01", "quantity": 0.5}], "2026-10", "小时", 200)
        self.assertEqual(
            daily_rows([{"date": "2026-10-01", "quantity": 20.5}], "2026-10", "小时", 200)[0]["amount"], 4100
        )
        with self.assertRaises(PermissionError):
            self.save(actor=None)

    def test_frozen_source_versions_and_distinct_feeding_workflow(self):
        item = self.save(actor=STAFF)
        preview = settlement.preview(self.conn, ADMIN, {"month": "2026-10", "pi": "纪老师"})
        self.save(
            self.sheet(rows=[{"date": "2026-10-01", "quantity": 20}]), identity=item["id"], expected=item["updatedAt"]
        )
        with self.assertRaises(StaleWriteError), self.conn:
            settlement.initiate(
                self.conn, ADMIN, {"month": "2026-10", "pi": "纪老师", "sourceVersions": preview["sourceVersions"]}
            )
        result = self.start()
        with self.assertRaises(ValueError):
            self.save(self.sheet(), identity=item["id"], expected=repo.get(self.conn, item["id"])[1]["updatedAt"])
        with self.conn:
            feeding = {**result["statement"], "sourceType": "pi_merged_quantity_sheet", "totalAmount": 55}
            workflow, *_ = save_billing_statement_workflow(self.conn, feeding, [], ADMIN)
        self.assertNotEqual(workflow["id"], result["workflow"]["id"])
        self.assertEqual(self.start()["statement"]["totalAmount"], 600)
        self.assertTrue(result["statement"]["documentNumber"].startswith("CL-PTS-"))

    def test_shared_workflow_permissions_registration_lock_and_unlock(self):
        self.save(actor=STAFF)
        result = self.start()
        identity = result["workflow"]["id"]
        self.assertTrue(withdrawal_capabilities(self.conn, STAFF, [identity])[identity])
        outsider = {**STAFF, "id": "outsider"}
        self.assertFalse(withdrawal_capabilities(self.conn, outsider, [identity])[identity])
        operator = {**outsider, "billingLockAllowed": True}
        self.assertTrue(withdrawal_capabilities(self.conn, operator, [identity])[identity])
        require_workflow_withdrawal(self.conn, operator, result["workflow"], to_status="statement_generated")
        with self.assertRaises(PermissionError):
            require_workflow_withdrawal(self.conn, outsider, result["workflow"], to_status="statement_generated")
        with self.conn:
            workflow, _, _ = update_workflow_status(
                self.conn,
                identity,
                "statement_archived",
                STAFF,
                registration={
                    "signedStatementReturned": True,
                    "reimbursementFormReturned": True,
                    "reimbursementForms": [{"formNo": "TS-001", "amount": 300, "fundingBookNo": "F001"}],
                },
            )
            self.assertEqual(workflow["workflowStatus"], "statement_archived")
            update_workflow_status(self.conn, identity, "statement_locked", ADMIN)
        self.assertFalse(withdrawal_capabilities(self.conn, STAFF, [identity])[identity])
        with self.conn:
            workflow, _, _ = update_workflow_status(self.conn, identity, "statement_archived", ADMIN)
            self.assertEqual(workflow["workflowStatus"], "statement_archived")
            update_workflow_status(self.conn, identity, "statement_sent", STAFF, note="修正登记")
            update_workflow_status(self.conn, identity, "statement_generated", STAFF, note="修正数量")
        self.assertNotIn("service-1", repo.protected_ids(self.conn))

    def test_price_snapshot_and_summary_facility_only_service(self):
        first = self.save()
        self.save(self.sheet(id="island", facility="bioisland", customerType="external"))
        result = self.start()
        self.assertEqual(result["statement"]["totalAmount"], 700)
        summary = merge_service_rows([], settlement.export_sheets(self.conn, "2026-10"), {}, {})
        self.assertEqual(
            {r["facility"]: r["technicalServiceAmount"] for r in summary}, {"珠江新城": 300, "生物岛": 400}
        )
        workbook = load_workbook(io.BytesIO(build_monthly_summary_xlsx("2026-10", summary)), data_only=False)
        self.assertEqual(workbook.active["S2"].value, "技术服务费（元）")
        self.assertEqual(sum(workbook.active.cell(i, 19).value for i in (3, 4)), 700)
        self.assertEqual(workbook.active["S3"].number_format, "0.00")
        self.assertEqual(first["priceSnapshot"]["version"], "eye-research-2024-4-v1")
        self.assertEqual(result["statement"]["serviceDayCount"], 4)

    def test_pdf_total_all_pages_and_shared_ledger(self):
        self.save(self.sheet(rows=[{"date": f"2026-10-{i:02}", "quantity": 1} for i in range(1, 26)]))
        result = self.start()
        html = billing_statement_html(result["statement"], result["lines"])
        self.assertEqual(html.count('class="technical-service-print-page"'), 2)
        self.assertIn("750.00 元", html)
        self.assertIn("第 2 / 2 页", html)
        self.assertIn("2026-10-25", html)
        obligations = list_obligations(self.conn, ADMIN, {})["items"]
        self.assertEqual(obligations[0]["payableAmount"], 750)
        self.assertEqual(obligations[0]["feeType"], "technical_service")

    def test_saved_price_and_project_snapshot_are_not_repriced(self):
        item = self.save(self.sheet(pi="伪造负责人", funding="伪造经费"))
        self.assertEqual(item["pi"], "纪老师")
        self.assertEqual(item["funding"], "项目来源")
        changed = {**item["priceSnapshot"], "internalPrice": 99}
        self.conn.execute(
            "UPDATE technical_service_catalog SET payload=? WHERE id=?", (json.dumps(changed), item["serviceId"])
        )
        self.conn.commit()
        revised = self.save(
            self.sheet(rows=[{"date": "2026-10-01", "quantity": 1}]), identity=item["id"], expected=item["updatedAt"]
        )
        self.assertEqual(revised["totalAmount"], 30)
        self.assertEqual(revised["priceSnapshot"]["internalPrice"], 30)

    def test_candidate_scope_paging_uses_full_pi_total_and_stable_count(self):
        self.save()
        self.save(self.sheet(id="island", facility="bioisland"))
        application = {"iacuc": "Z2026002", "pi": "另一位负责人", "project": "另一项目", "funding": "另一经费"}
        self.conn.execute(
            "INSERT INTO experiment_applications(id,iacuc,imported_at,payload) VALUES(?,?,?,?)",
            ("app2", "Z2026002", "2026-10-01", json.dumps(application)),
        )
        self.conn.commit()
        invalidate_data_cache()
        invalidate_data_cache_prefixes("applications_by_iacuc")
        self.save(self.sheet(id="another", iacuc="Z2026002"))
        result = settlement.candidates(self.conn, ADMIN, {"month": "2026-10", "limit": 1, "offset": 0})
        self.assertEqual(result["page"]["total"], 2)
        self.assertTrue(result["page"]["hasMore"])
        searched = settlement.candidates(self.conn, ADMIN, {"month": "2026-10", "keyword": "Z2026001"})
        self.assertEqual(searched["items"][0]["totalAmount"], 600)

    def test_shared_claim_allocates_to_separate_fee_obligations_without_collision(self):
        self.save()
        result = self.start()
        with self.conn:
            feeding = {**result["statement"], "sourceType": "pi_merged_quantity_sheet", "totalAmount": 55}
            save_billing_statement_workflow(self.conn, feeding, [], ADMIN)
        obligations = list_obligations(self.conn, ADMIN, {})["items"]
        self.assertEqual({o["feeType"] for o in obligations}, {"feeding", "technical_service"})
        self.assertEqual(len({o["workflowId"] for o in obligations}), 2)
        claim = save_claim(
            self.conn,
            ADMIN,
            None,
            {
                "documentNumber": "SHARED-TS-FEED",
                "fundingLines": [{"fundBookNo": "F1", "fundingOwner": "纪老师", "reimbursementAmount": 355}],
            },
        )["item"]
        line = claim["fundingLines"][0]
        for obligation in obligations:
            allocation = create_allocation(
                self.conn,
                ADMIN,
                {
                    "claimId": claim["id"],
                    "fundingLineId": line["id"],
                    "obligationId": obligation["id"],
                    "amount": obligation["payableAmount"],
                },
            )["item"]
            confirm_allocation(self.conn, ADMIN, allocation["id"])
        items = list_obligations(self.conn, ADMIN, {})["items"]
        self.assertEqual(sum(o["allocatedAmount"] for o in items), 355)
        self.assertTrue(all(o["outstandingAmount"] == 0 for o in items))

    def test_shared_reimbursement_form_not_counted_twice(self):
        form = {"formNo": "F01", "fundingBookNo": "B01", "amount": 300}
        self.assertEqual(
            build_pi_reimbursement_context([{"reimbursementForms": [form]}, {"reimbursementForms": [form]}])[
                "reimbursementAmount"
            ],
            300,
        )


if __name__ == "__main__":
    unittest.main()
